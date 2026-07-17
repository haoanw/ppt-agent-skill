#!/usr/bin/env python3
"""Rebuild the FMC3 product manual from a blank dark-tech canvas."""

from __future__ import annotations

import argparse
import html
import json
import math
import posixpath
import re
import shutil
import subprocess
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from zipfile import ZipFile

from lxml import etree
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

import build_product_manual_dark_tech as workflow


ROOT = Path(__file__).resolve().parents[2]
BD = ROOT / "products/BD"
SOURCE = BD / "FMC³ Robotics_product Manual_0606.pptx"
DEFAULT_OUTPUT = BD / "FMC3_Robotics_Product_Manual_dark_tech_redesigned"
BASENAME = "FMC3_Robotics_Product_Manual_dark_tech_redesigned"
WHITE_LOGO = (
    BD
    / "FMC3_C5_Brochure_Deutsch_Redesigned_2_HighRes_Vector_dark_tech"
    / "images/image3.png"
)

W = 13.333
H = 7.5
BG = RGBColor(0x04, 0x0A, 0x1B)
BG_2 = RGBColor(0x08, 0x17, 0x30)
PANEL = RGBColor(0x0B, 0x1B, 0x36)
PANEL_2 = RGBColor(0x10, 0x28, 0x49)
GRID = RGBColor(0x12, 0x2A, 0x49)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
INDIGO = RGBColor(0x63, 0x66, 0xF1)
YELLOW = RGBColor(0xFD, 0xE0, 0x47)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0xC9, 0xD4, 0xE3)
MUTED = RGBColor(0x87, 0x9A, 0xB3)
LIGHT = RGBColor(0xF3, 0xF8, 0xFC)
RED = RGBColor(0xF4, 0x3F, 0x5E)
GREEN = RGBColor(0x10, 0xB9, 0x81)
FONT = "Aptos"
MONO = "Aptos Mono"


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value, encoding="utf-8")


def write_json(path: Path, value: object) -> None:
    write_text(path, json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def extract_assets(output_dir: Path) -> dict[int, list[Path]]:
    # Workflow planning contracts resolve supplied visuals under output/images/.
    assets_dir = output_dir / "images"
    assets_dir.mkdir(parents=True, exist_ok=True)
    mapping: dict[int, list[Path]] = {}
    with ZipFile(SOURCE) as archive:
        for page in range(1, 31):
            rel_path = f"ppt/slides/_rels/slide{page}.xml.rels"
            root = etree.fromstring(archive.read(rel_path))
            seen: set[str] = set()
            files: list[Path] = []
            for node in root:
                target = node.get("Target", "")
                if "/media/" not in target:
                    continue
                media = posixpath.normpath(posixpath.join("ppt/slides", target))
                if media in seen:
                    continue
                seen.add(media)
                ext = PurePosixPath(media).suffix.lower()
                if ext not in {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff"}:
                    continue
                destination = assets_dir / f"slide-{page:02d}-asset-{len(files) + 1:02d}{ext}"
                destination.write_bytes(archive.read(media))
                files.append(destination)
            mapping[page] = files
    return mapping


def is_raster(path: Path) -> bool:
    return path.suffix.lower() in {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff"}


def raster_assets(mapping: dict[int, list[Path]], page: int) -> list[Path]:
    return [path for path in mapping[page] if is_raster(path)]


def tint_map(source: Path, destination: Path) -> Path:
    with Image.open(source).convert("RGBA") as image:
        pixels = image.load()
        for y in range(image.height):
            for x in range(image.width):
                r, g, b, a = pixels[x, y]
                darkness = 255 - int((r + g + b) / 3)
                if darkness < 18:
                    pixels[x, y] = (0, 0, 0, 0)
                else:
                    alpha = min(220, int(darkness * 1.55))
                    pixels[x, y] = (34, 211, 238, alpha)
        bbox = image.getbbox()
        if bbox:
            image = image.crop(bbox)
        image.save(destination)
    return destination


def crop_cover(source: Path, destination: Path, ratio: float = 16 / 9) -> Path:
    with Image.open(source).convert("RGB") as image:
        current = image.width / image.height
        if current > ratio:
            width = int(image.height * ratio)
            left = (image.width - width) // 2
            image = image.crop((left, 0, left + width, image.height))
        else:
            height = int(image.width / ratio)
            top = (image.height - height) // 2
            image = image.crop((0, top, image.width, top + height))
        image.resize((1920, 1080), Image.Resampling.LANCZOS).save(destination, quality=94)
    return destination


def set_alpha(shape, transparency: int) -> None:
    solid_fill = shape.fill._xPr.solidFill
    color_nodes = list(solid_fill)
    if not color_nodes:
        return
    color = color_nodes[0]
    for child in list(color):
        if child.tag.endswith("alpha"):
            color.remove(child)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(max(0, min(100, 100 - transparency)) * 1000))
    color.append(alpha)


def rect(slide, x, y, w, h, fill=PANEL, line=GRID, radius=False, transparency=0):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if transparency:
        set_alpha(shape, transparency)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.7)
    return shape


def line(slide, x1, y1, x2, y2, color=GRID, width=0.7):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def add_text(
    slide,
    items,
    x,
    y,
    w,
    h,
    *,
    size=12,
    color=TEXT,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margins=(0.03, 0.03, 0.03, 0.03),
    spacing=1.05,
    bullet=False,
    fit=False,
):
    values = [items] if isinstance(items, str) else list(items)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margins[0])
    frame.margin_top = Inches(margins[1])
    frame.margin_right = Inches(margins[2])
    frame.margin_bottom = Inches(margins[3])
    frame.vertical_anchor = valign
    frame.word_wrap = True
    if fit:
        frame.fit_text(font_family=font, max_size=int(size))
    for index, value in enumerate(values):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = value
        paragraph.alignment = align
        paragraph.space_after = Pt(max(1, size * 0.22))
        paragraph.line_spacing = spacing
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        if bullet:
            ppr = paragraph._p.get_or_add_pPr()
            bu = OxmlElement("a:buChar")
            bu.set("char", "•")
            ppr.append(bu)
            paragraph.margin_left = Inches(0.18)
            paragraph.indent = Inches(-0.10)
    return box


def add_rich_lines(slide, lines_data, x, y, w, h, size=11):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for index, (label, value) in enumerate(lines_data):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_after = Pt(4); p.line_spacing = 1.0
        r = p.add_run(); r.text = label
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE
        if value:
            r = p.add_run(); r.text = f"  {value}"
            r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = TEXT
    return box


def add_picture_contain(slide, path: Path, x, y, w, h, padding=0.0):
    if not path.is_file() or not is_raster(path):
        return None
    try:
        with Image.open(path) as image:
            iw, ih = image.size
    except Exception:
        return None
    w2, h2 = max(0.05, w - 2 * padding), max(0.05, h - 2 * padding)
    scale = min(w2 / iw, h2 / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(
        str(path), Inches(x + padding + (w2 - pw) / 2), Inches(y + padding + (h2 - ph) / 2),
        width=Inches(pw), height=Inches(ph),
    )


def add_logo(slide):
    if WHITE_LOGO.is_file():
        add_picture_contain(slide, WHITE_LOGO, 11.55, 0.22, 1.18, 0.52)


def add_background(slide, page: int, source_text: list[str], *, section="FMC³ ROBOTICS") -> None:
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = BG
    for x in (0.8, 2.6, 4.4, 6.2, 8.0, 9.8, 11.6):
        line(slide, x, 0, x, H, GRID, 0.35)
    for y in (1.25, 2.75, 4.25, 5.75, 7.1):
        line(slide, 0, y, W, y, GRID, 0.35)
    line(slide, 0.45, 0.18, 1.15, 0.18, CYAN, 1.1)
    line(slide, 0.45, 0.18, 0.45, 0.52, CYAN, 1.1)
    line(slide, 0.55, 7.16, 12.78, 7.16, GRID, 0.8)
    add_logo(slide)
    url = next((item for item in source_text if item.lower().startswith("www.")), "www.fmc3robotics.ai")
    date = next((item for item in source_text if re.fullmatch(r"\d{2}/\d{2}/\d{4}|\d{4}-\d{2}", item)), "")
    add_text(slide, section, 0.55, 7.21, 3.0, 0.18, size=7.5, color=MUTED, bold=True, font=MONO)
    add_text(slide, url, 4.75, 7.20, 3.8, 0.18, size=7.5, color=MUTED, align=PP_ALIGN.CENTER, font=MONO)
    if date:
        add_text(slide, date, 9.5, 7.20, 1.4, 0.18, size=7.5, color=MUTED, align=PP_ALIGN.RIGHT, font=MONO)
    add_text(slide, f"{page:02d} / 30", 11.55, 7.18, 1.15, 0.2, size=8.5, color=CYAN, bold=True, align=PP_ALIGN.RIGHT, font=MONO)


def add_title(slide, title: str, *, eyebrow: str | None = None, subtitle: str | None = None, compact=False):
    if eyebrow:
        parts = eyebrow.split(" · ")
        x = 0.58
        for index, part in enumerate(parts):
            width = max(0.72, min(4.8, 0.10 * len(part) + 0.38))
            add_text(slide, part, x, 0.34, width, 0.24, size=8.5, color=CYAN if index == 0 else MUTED, bold=True, font=MONO)
            x += width + 0.18
    top = 0.56 if eyebrow else 0.38
    size = 23 if compact else 28
    height = 0.62 if compact else 0.76
    add_text(slide, title, 0.55, top, 10.6, height, size=size, color=WHITE, bold=True, spacing=0.95)
    if subtitle:
        add_text(slide, subtitle, 0.58, top + height + 0.02, 10.4, 0.36, size=11.5, color=TEXT)


def add_section_label(slide, text: str, x, y, w, *, accent=CYAN):
    line(slide, x, y + 0.32, x + w, y + 0.32, accent, 0.8)
    add_text(slide, text, x, y, w, 0.29, size=10, color=accent, bold=True, font=MONO)


def add_card(slide, title: str, body: list[str], x, y, w, h, *, accent=CYAN, body_size=10.5, title_size=10.5, bullet=False):
    rect(slide, x, y, w, h, PANEL, GRID)
    rect(slide, x, y, 0.035, h, accent, None)
    add_text(slide, title, x + 0.18, y + 0.12, w - 0.32, 0.3, size=title_size, color=accent, bold=True)
    add_text(slide, body, x + 0.18, y + 0.50, w - 0.32, h - 0.58, size=body_size, color=TEXT, spacing=1.0, bullet=bullet)


def add_metric(slide, value: str, label: str, x, y, w, *, accent=CYAN):
    rect(slide, x, y, w, 0.78, PANEL, GRID)
    add_text(slide, value, x + 0.15, y + 0.10, w - 0.3, 0.36, size=18, color=WHITE, bold=True)
    add_text(slide, label, x + 0.15, y + 0.49, w - 0.3, 0.18, size=7.2, color=accent, bold=True, font=MONO)


def split_sections(items: list[str], headings: list[str]) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    title = ""
    body: list[str] = []
    heading_set = set(headings)
    for item in items:
        if item in heading_set:
            if title or body:
                sections.append((title, body))
            title, body = item, []
        else:
            body.append(item)
    if title or body:
        sections.append((title, body))
    return sections


def pairs(items: list[str]) -> list[tuple[str, str]]:
    result = []
    for index in range(0, len(items), 2):
        result.append((items[index], items[index + 1] if index + 1 < len(items) else ""))
    return result


def text_without(source: list[str], removals: list[str]) -> list[str]:
    pending = list(source)
    for value in removals:
        if value in pending:
            pending.remove(value)
    return [item for item in pending if not re.fullmatch(r"\d+", item)]


def clean_footer_items(items: list[str]) -> list[str]:
    return [
        item for item in items
        if not item.lower().startswith("www.")
        and not re.fullmatch(r"\d{2}/\d{2}/\d{4}|\d{4}-\d{2}", item)
        and not re.fullmatch(r"\d+", item)
    ]


def set_notes(slide, notes: list[str]) -> None:
    if not notes:
        return
    try:
        slide.notes_slide.notes_text_frame.text = "\n".join(notes)
    except Exception:
        pass


def build_cover(prs, texts, assets, runtime):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    hero = crop_cover(assets[1], runtime / "cover-hero.jpg")
    slide.shapes.add_picture(str(hero), 0, 0, width=Inches(W), height=Inches(H))
    rect(slide, 0, 0, 7.0, H, BG, None, transparency=12)
    line(slide, 0.62, 0.55, 2.0, 0.55, CYAN, 1.5)
    add_text(slide, "FMC³ Robotics", 0.62, 1.08, 7.4, 1.06, size=43, color=WHITE, bold=True)
    add_text(slide, "The Embodied AI Brain from Europe", 0.66, 2.32, 6.2, 0.55, size=19, color=CYAN, bold=True)
    add_text(slide, "Product Manual", 0.66, 4.95, 2.4, 0.34, size=11, color=WHITE, bold=True, font=MONO)
    add_text(slide, ["Mar, 2026", "Confidential"], 0.66, 5.36, 2.4, 0.75, size=10.5, color=TEXT, font=MONO)
    add_logo(slide)
    add_text(slide, "01 / 30", 11.55, 7.18, 1.15, 0.2, size=8.5, color=CYAN, bold=True, align=PP_ALIGN.RIGHT, font=MONO)
    return slide


def build_positioning(prs, texts, assets, runtime):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 2, texts, section="POSITIONING")
    add_title(slide, "Positioning")
    strategy = [texts[17], texts[18]]
    add_text(slide, strategy, 0.62, 1.15, 6.0, 1.28, size=13.5, color=WHITE, bold=True, spacing=1.02, bullet=True)
    map_path = tint_map(assets[0], runtime / "world-map-cyan.png")
    add_picture_contain(slide, map_path, 7.15, 1.0, 5.5, 4.55)
    line(slide, 8.15, 3.2, 10.9, 3.85, CYAN, 1.2)
    rect(slide, 7.86, 2.98, 0.16, 0.16, CYAN, None, radius=True)
    rect(slide, 10.82, 3.77, 0.16, 0.16, BLUE, None, radius=True)
    europe = [texts[2], texts[3], texts[4], texts[8], texts[9]]
    china = [texts[5], texts[6], texts[7], texts[10], texts[11]]
    add_card(slide, "Europe HQ", europe, 0.62, 2.7, 5.8, 1.55, accent=CYAN, body_size=10.2)
    add_card(slide, "China Office", china, 0.62, 4.45, 5.8, 1.55, accent=BLUE, body_size=10.2)
    add_text(slide, ["HQ", "Product &Engineering"], 7.15, 4.9, 2.15, 0.72, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, ["China office", "RD&Product design"], 10.16, 4.9, 2.25, 0.72, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return slide


def build_brain(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 3, texts, section="FMC³ BRAIN")
    add_title(slide, "One Brain filt all", subtitle="Data loop, embodied intelligence and multi-body adaptation")
    items = clean_footer_items(text_without(texts, ["One Brain filt all"]))
    left = items[:3]
    add_picture_contain(slide, assets[2], 4.35, 1.55, 4.55, 3.15)
    for i, value in enumerate(left):
        add_card(slide, value, [], 0.58, 1.55 + i * 1.2, 3.2, 0.88, accent=CYAN, title_size=12)
        line(slide, 3.78, 1.99 + i * 1.2, 4.4, 2.3 + i * 0.55, CYAN, 0.8)
    right = [(items[3], items[4]), (items[5], items[6]), (items[7], items[8])]
    for i, (heading, detail) in enumerate(right):
        add_card(slide, heading, [detail], 9.45, 1.55 + i * 1.2, 3.2, 0.88, accent=BLUE, title_size=10.5, body_size=8.4)
    add_picture_contain(slide, assets[12], 5.15, 4.68, 3.0, 1.2)
    add_text(slide, items[9:], 8.75, 4.72, 3.85, 1.60, size=8.6, color=TEXT, bullet=True)
    for index, asset in enumerate(assets[13:17]):
        add_picture_contain(slide, asset, 0.65 + index * 2.0, 5.78, 1.55, 0.95)
    return slide


def build_team(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 4, texts, section="CORE TEAM")
    add_title(slide, "Core team", subtitle="A Team of European Intelligent Systems Scientists")
    people = [
        ("Dr. Wang Cheng", "CEO", texts[9:13], assets[3]),
        ("Prof. Dr. Giesler Bjoern", "CTO", texts[13:16], assets[4]),
        ("Fan David, PhD", "COO", texts[16:19], assets[1]),
        ("Dr. Huang Dong", "CPO", texts[19:22], assets[2]),
        ("Obermeier Florian", "MD Germany", texts[25:28], assets[5]),
    ]
    for i, (name, role, bio, image) in enumerate(people):
        x = 0.56 + i * 2.52
        rect(slide, x, 1.55, 2.30, 5.15, PANEL, GRID)
        rect(slide, x, 1.55, 2.30, 0.045, CYAN if i < 4 else BLUE, None)
        add_picture_contain(slide, image, x + 0.10, 1.68, 2.10, 1.55)
        add_text(slide, name, x + 0.14, 3.34, 2.02, 0.42, size=11.1, color=WHITE, bold=True)
        add_text(slide, role, x + 0.14, 3.77, 2.02, 0.24, size=8.2, color=CYAN, bold=True, font=MONO)
        add_text(slide, bio, x + 0.14, 4.12, 2.02, 2.35, size=8.0, color=TEXT, spacing=0.98, bullet=True)
    return slide


def build_gdpr(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 5, texts, section="GDPR / EU AI ACT")
    add_title(slide, "GDPR Data Compliance", subtitle="Localized Data Closed-Loop Architecture (Germany/Europe)")
    body = clean_footer_items(text_without(texts, ["GDPR Data Compliance", "Localized Data Closed-Loop Architecture (Germany/Europe)"]))
    add_card(slide, body[0], body[1:2], 0.62, 1.48, 3.0, 1.15, accent=RED, body_size=9.4)
    add_picture_contain(slide, assets[0], 0.78, 1.68, 0.42, 0.42)
    add_section_label(slide, body[2], 4.0, 1.12, 8.65)
    stages = [
        (body[3], body[4], body[5], assets[2]),
        (body[6], body[7], body[8], assets[4]),
        (body[9], body[10], body[11], assets[5]),
    ]
    for i, (title, description, state, icon) in enumerate(stages):
        y = 1.48 + i * 1.5
        rect(slide, 4.0, y, 8.65, 1.10, PANEL, GRID)
        add_picture_contain(slide, icon, 4.22, y + 0.25, 0.52, 0.52)
        add_text(slide, title, 4.92, y + 0.16, 2.4, 0.25, size=11, color=WHITE, bold=True)
        add_text(slide, description, 4.92, y + 0.48, 5.35, 0.39, size=9.1, color=TEXT)
        rect(slide, 10.62, y + 0.28, 1.58, 0.48, PANEL_2, CYAN, radius=True)
        add_text(slide, state, 10.7, y + 0.35, 1.42, 0.22, size=8.2, color=CYAN, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    principles = body[-4:]
    for i, value in enumerate(principles):
        x = 4.0 + i * 2.16
        rect(slide, x, 6.10, 2.02, 0.48, BG_2, GRID)
        add_text(slide, value, x + 0.08, 6.18, 1.86, 0.25, size=7.8, color=TEXT, align=PP_ALIGN.CENTER)
    add_text(slide, "Fermi Robot Business Plan | Confidential", 0.66, 6.70, 3.1, 0.18, size=7.2, color=MUTED, font=MONO)
    return slide


def build_safety(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 6, texts, section="EU SAFETY")
    title = 'Entering the European Industry: Building a Moat Based on "Deterministic, Safe, and Compliant" Standards'
    add_title(slide, title, compact=True)
    body = clean_footer_items(text_without(texts, [title, "The Four Hard Requirements for Embodied AI Robots under the \"Machinery Regulation\""]))
    add_picture_contain(slide, assets[1], 0.62, 1.30, 4.05, 2.2)
    add_card(slide, "Collaborative Safety & Compliance", body[1:7], 4.95, 1.30, 7.68, 2.2, accent=CYAN, body_size=9.7)
    headings = [
        "1. Safety of Control Systems", "2.Special Requirements for Autonomous Mobility",
        "3.Risk Monitoring & Cybersecurity", "4. Software Safety & Self-learning",
    ]
    sections = split_sections(body[7:], headings)
    icons = assets[2:6]
    for i, (heading, content) in enumerate(sections[:4]):
        x = 0.62 + i * 3.02
        add_card(slide, heading, content, x, 4.0, 2.78, 2.25, accent=CYAN if i < 2 else BLUE, body_size=8.9, title_size=9.2)
        if i < len(icons):
            add_picture_contain(slide, icons[i], x + 2.05, 4.15, 0.52, 0.52)
        add_text(slide, str(i + 1), x + 0.12, 5.82, 0.35, 0.25, size=12, color=CYAN, bold=True, font=MONO)
    add_text(slide, 'The Four Hard Requirements for Embodied AI Robots under the "Machinery Regulation"', 0.66, 6.48, 11.9, 0.34, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    return slide


def build_partners(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 7, texts, section="ECOSYSTEM")
    add_title(slide, "Our Partners and Clients")
    logos = [assets[i] for i in [0, 1, 2, 3, 4, 5, 6, 7, 8, 11, 12]]
    layout = [(0.7, 1.35), (3.75, 1.35), (6.80, 1.35), (9.85, 1.35),
              (1.85, 3.13), (4.90, 3.13), (7.95, 3.13),
              (0.95, 4.91), (4.0, 4.91), (7.05, 4.91), (10.1, 4.91)]
    for path, (x, y) in zip(logos, layout):
        rect(slide, x, y, 2.55, 1.25, LIGHT, CYAN)
        add_picture_contain(slide, path, x + 0.12, y + 0.12, 2.31, 1.01)
    return slide


def build_portfolio(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 8, texts, section="PRODUCT PORTFOLIO")
    add_title(slide, "FMC³ Brain+ODM", eyebrow="Products")
    body = clean_footer_items(text_without(texts, ["Products", "FMC³ Brain+ODM"]))
    footnote = body[-1]
    brain = assets[8]
    rect(slide, 4.75, 2.42, 3.85, 1.52, PANEL_2, CYAN)
    add_picture_contain(slide, brain, 5.10, 2.55, 3.15, 0.95)
    add_text(slide, "FMC³ Brain + FMC³ Data Acquisition Suit", 4.92, 3.55, 3.50, 0.26, size=9.4, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    product_assets = [assets[10], assets[3], assets[5], assets[4], assets[12], assets[11], assets[6]]
    label_groups = [
        ["Robo dog"], ["Commercial & Industrial Cleaning"], ["Workbot"], ["Cobot"],
        ["Industrial", "Service Robot"], ["Humanoid Home Service Robot"], ["GreetingBot"],
    ]
    positions = [(0.55, 1.55), (2.25, 1.55), (9.25, 1.55), (10.95, 1.55), (0.55, 4.32), (2.25, 4.32), (9.75, 4.32)]
    for labels, image, (x, y) in zip(label_groups, product_assets, positions):
        rect(slide, x, y, 1.55, 1.78, PANEL, GRID)
        add_picture_contain(slide, image, x + 0.12, y + 0.10, 1.31, 1.12)
        add_text(slide, labels, x + 0.08, y + 1.27, 1.39, 0.43, size=7.8, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        line(slide, x + 0.78, y + (1.78 if y < 3 else 0), 6.67, 3.18, GRID, 0.6)
    add_text(slide, footnote, 3.10, 6.25, 7.15, 0.48, size=8.4, color=MUTED, align=PP_ALIGN.CENTER)
    return slide


def build_brain_suite(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 9, texts, section="BRAIN + DATA")
    add_title(slide, "FMC³ Brain+Data Acquisition Suit", eyebrow="Products")
    body = clean_footer_items(text_without(texts, ["Products", "FMC³ Brain+Data Acquisition Suit"]))
    add_picture_contain(slide, assets[5], 0.72, 1.55, 3.05, 1.65)
    metrics = body[13:21]
    for i, (label, value) in enumerate(pairs(metrics)):
        y = 3.30 + i * 0.62
        rect(slide, 0.62, y, 3.15, 0.51, PANEL if i % 2 == 0 else BG_2, GRID)
        add_text(slide, label, 0.76, y + 0.09, 1.48, 0.22, size=7.5, color=CYAN, bold=True, font=MONO)
        add_text(slide, value, 2.30, y + 0.08, 1.30, 0.24, size=9.0, color=WHITE, bold=True, align=PP_ALIGN.RIGHT)
    add_picture_contain(slide, assets[13], 4.18, 1.55, 3.25, 2.0)
    add_text(slide, body[21:], 4.22, 3.62, 3.18, 2.35, size=9.2, color=TEXT, bullet=True)
    headings = ["Realtime Teleoperation & Emergency Takeover", "FullLink Data Toolchain", "Core Data Set"]
    sections = split_sections(body[:13], headings)
    for i, (heading, content) in enumerate(sections):
        y = [1.55, 3.18, 4.81][i]
        height = 1.40 if i < 2 else 1.62
        add_card(slide, heading, content, 7.82, y, 4.80, height, accent=CYAN if i == 0 else BLUE, body_size=8.4 if i < 2 else 7.8)
    return slide


def build_omnigrab(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 10, texts, section="DATA REFINERY")
    title = "FMC³ Omni-Grab : Ultra-Fidelity, Body-Free Human-Ego-Centric Data Capture System"
    add_title(slide, title, eyebrow="Products · FMC³ Brain+Data Acquisition Suit", compact=True)
    body = clean_footer_items(text_without(texts, [title, "Products", "FMC³ Brain+Data Acquisition Suit"]))
    hardware_heads = ["Active IR Stereo RGB-D Headset", "400-Point High-Density Haptic Glove", "Portable Data Sync Hub"]
    sections = split_sections(body, hardware_heads + ["End-to-End Automated Processing Pipeline", "DATA REFINERY PIPELINE"])
    hardware = sections[:3]
    hardware_images = [assets[2], assets[3], assets[4]]
    for i, ((heading, content), image) in enumerate(zip(hardware, hardware_images)):
        x = 0.60 + i * 2.18
        rect(slide, x, 1.35, 1.95, 5.30, PANEL, GRID)
        add_picture_contain(slide, image, x + 0.16, 1.48, 1.63, 0.95)
        add_text(slide, heading, x + 0.14, 2.45, 1.67, 0.42, size=9.2, color=CYAN, bold=True)
        add_text(slide, content, x + 0.14, 2.92, 1.67, 3.45, size=7.8, color=TEXT)
    pipeline_items = []
    for _, content in sections[3:]:
        pipeline_items.extend(content)
    step_heads = [item for item in pipeline_items if item.startswith("STEP ")]
    descriptions = [item for item in pipeline_items if not item.startswith("STEP ")]
    add_section_label(slide, "End-to-End Automated Processing Pipeline", 7.20, 1.05, 5.45)
    add_section_label(slide, "DATA REFINERY PIPELINE", 7.20, 1.34, 5.45)
    step_titles = [
        "High-Precision Synchronized Capture:", "Automatic Scene Segmentation & Slicing:",
        "Smart Cleaning & Multi-Dimensional QA:", "Auto-Annotation & Programmatic Augmentation:",
        "Standardized Packaging & Universal Export:",
    ]
    cursor = 0
    for i, heading in enumerate(step_titles):
        desc = []
        if heading in descriptions:
            idx = descriptions.index(heading)
            if idx + 1 < len(descriptions): desc = [descriptions[idx + 1]]
        y = 1.72 + i * 0.98
        add_text(slide, step_heads[i] if i < len(step_heads) else f"STEP {i+1}", 7.20, y, 0.72, 0.25, size=8, color=CYAN, bold=True, font=MONO)
        add_text(slide, heading, 8.0, y, 4.55, 0.25, size=9, color=WHITE, bold=True)
        add_text(slide, desc, 8.0, y + 0.29, 4.55, 0.48, size=7.8, color=TEXT)
        line(slide, 7.55, y + 0.28, 7.55, y + 0.92, GRID, 0.8)
        cursor += 1
    return slide


PRODUCT_CONFIG = {
    11: ("Robo dogD1 MAX PRO", ["Highlights", "Ultrahigh Payload", "Reliable Battery Life", "Outstanding Mobility Performance", "Superior Environmental Adaptability"], 1),
    12: ("Robo dogD1 MAX PRO", ["Functions", "Parameters", "Performance Parameters", "Robot Joint Parameters", "Sensing Unit"], 1),
    13: ("Robo dogD1 MAX", ["Highlights", "1. Core Performance Advantages", "2. Intelligent Locomotion & Terrain Adaptability", "3. FullScene Perception & Navigation", "4. MultiScenario Application Capabilities"], 1),
    14: ("Robo dogD1 MAX", ["Basic Information", "Performance Parameters", "Robot Joint Parameters", "Sensing Unit"], 1),
    15: ("Robo dog D1 Ultra", ["Highlights", "Robot Joint Parameters", "Sensors", "Performance Parameters"], 1),
    16: ("Robo dog D1 Pro", ["Highlights"], 1),
    17: ("Robo dog", ["Functions", "Parameters"], 1),
    18: ("Commercial & Industrial Cleaning", ["Functions", "Parameters", "Application Scenarios", "Applicable Floor Types"], None),
    20: ("Workbot", ["Functions", "Parameters", "Application Scenarios"], 2),
    25: ("Cobot", ["Functions", "Parameters", "Application Scenarios"], 1),
    26: ("Industrial Service Robot", ["Functions", "Parameters", "Application Scenarios"], 0),
    27: ("Humanoid Home Service Robot", ["Functions", "Parameters", "Application Scenarios"], 0),
    28: ("GreetingBot Nova Service Robot", ["Application Scenarios", "Functions", "Parameters"], 2),
    29: ("IBen-A03 Service Robot", ["Application Scenarios", "Functions", "Parameters"], 2),
}


def product_sections(page: int, body: list[str], headings: list[str]) -> list[tuple[str, list[str]]]:
    if page == 11:
        return [
            ("Highlights", body[1:3]),
            (body[3], body[4:5]),
            (body[5], body[6:7]),
            (body[7], body[8:9]),
        ]
    if page == 12:
        return [
            ("Functions", body[0:5]),
            ("Parameters", body[7:16]),
            ("Performance Parameters", body[17:22]),
            ("Robot Joint Parameters", body[23:27]),
            ("Sensing Unit", body[28:31]),
        ]
    if page == 13:
        return [
            ("Highlights", body[1:7]),
            (body[7], body[8:13]),
            (body[13], body[14:19]),
            (body[19], body[20:27]),
        ]
    if page == 15:
        return [
            ("Highlights", body[0:7]),
            ("Robot Joint Parameters", body[9:13]),
            ("Sensors", body[14:16]),
            ("Performance Parameters", body[17:22]),
        ]
    if page == 16:
        return [
            ("Robo dog D1 Pro", body[0:2]),
            ("Highlights", body[3:14]),
        ]
    if page == 17:
        return [
            ("Functions", body[0:5] + body[6:9]),
            ("Parameters", body[9:32]),
        ]
    if page == 18:
        return [
            ("Functions", body[0:5]),
            ("Parameters", body[7:39]),
            ("Application Scenarios", body[39:45]),
        ]
    if page == 20:
        return [
            ("Functions", body[0:6]),
            ("Application Scenarios", body[8:13]),
            ("Parameters", body[14:38]),
        ]
    if page == 25:
        params = body[11:36]
        params.insert(params.index("Total DOF") + 1, "21")
        return [
            ("Functions", body[2:6]),
            ("Application Scenarios", body[7:11]),
            ("Parameters", params),
        ]
    if page == 26:
        return [
            ("Functions", body[2:8] + body[39:42]),
            ("Application Scenarios", body[9:15]),
            ("Parameters", body[15:39]),
        ]
    if page == 27:
        params = body[15:44]
        params.insert(params.index("Total DOF") + 1, "55")
        return [
            ("Functions", body[2:5]),
            ("Application Scenarios", body[6:15]),
            ("Parameters", params),
        ]
    return [(heading, content) for heading, content in split_sections(body, headings) if content]


def build_product_page(prs, page, texts, assets, cleaning_asset=None):
    name, headings, image_index = PRODUCT_CONFIG[page]
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, page, texts, section="PRODUCTS")
    add_title(slide, name, eyebrow="Products · FMC³ Brain+ODM", compact=True)
    removals = ["Products", "FMC³ Brain+ODM", name]
    if page == 13: removals.append("Robo dogD1 MAX")
    body = clean_footer_items(text_without(texts, removals))
    sections = product_sections(page, body, headings)
    image = cleaning_asset if page == 18 else assets[image_index] if image_index is not None else None
    rect(slide, 0.58, 1.30, 3.25, 5.72, PANEL, GRID)
    if page == 25:
        rect(slide, 0.72, 1.72, 2.95, 2.72, LIGHT, CYAN)
    if image:
        add_picture_contain(slide, image, 0.75, 1.47, 2.92, 3.05)
    # Use source scenario thumbnails when available.
    if page == 11:
        for i, thumb in enumerate(assets[2:5]):
            add_picture_contain(slide, thumb, 0.74 + (i % 2) * 1.45, 4.55 + (i // 2) * 0.88, 1.30, 0.76)
    elif page == 17 and len(assets) > 2:
        add_picture_contain(slide, assets[2], 0.72, 4.55, 2.95, 1.75)
    elif page == 18:
        if assets: add_picture_contain(slide, assets[0], 0.72, 4.70, 2.95, 1.05)
        if len(assets) > 2: add_picture_contain(slide, assets[2], 0.72, 5.88, 2.95, 0.72)
        add_text(slide, "Applicable Floor Types", 0.78, 6.58, 2.82, 0.22, size=8.2, color=CYAN, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    elif page == 20:
        for i, thumb in enumerate(assets[2:6]):
            add_picture_contain(slide, thumb, 0.72 + (i % 2) * 1.46, 4.38 + (i // 2) * 1.08, 1.30, 0.95)
    elif page == 25:
        for i, thumb in enumerate(assets[2:5]):
            add_picture_contain(slide, thumb, 0.72 + i * 0.95, 5.28, 0.78, 0.78)
    elif page == 27 and len(assets) > 2:
        add_picture_contain(slide, assets[1], 0.72, 4.75, 1.35, 1.15)
        add_picture_contain(slide, assets[2], 2.18, 4.75, 1.35, 1.15)
    content_x, content_w = 4.08, 8.55
    count = len(sections)
    if page == 12:
        lookup = {heading: content for heading, content in sections}
        add_card(slide, "Functions", lookup["Functions"], 4.08, 1.40, 4.12, 1.45, body_size=9.2)
        add_card(slide, "Performance Parameters", lookup["Performance Parameters"], 4.08, 3.03, 4.12, 1.55, body_size=8.8)
        add_card(slide, "Robot Joint Parameters", lookup["Robot Joint Parameters"], 4.08, 4.76, 4.12, 1.88, body_size=8.8)
        add_card(slide, "Parameters", lookup["Parameters"], 8.46, 1.40, 4.17, 2.35, accent=BLUE, body_size=8.2)
        add_card(slide, "Sensing Unit", lookup["Sensing Unit"], 8.46, 3.93, 4.17, 2.71, accent=BLUE, body_size=7.8)
    elif page == 15:
        lookup = {heading: content for heading, content in sections}
        add_card(slide, "Highlights", lookup["Highlights"], 4.08, 1.40, 4.12, 5.24, body_size=8.4)
        add_card(slide, "Robot Joint Parameters", lookup["Robot Joint Parameters"], 8.46, 1.40, 4.17, 1.54, accent=BLUE, body_size=8.8)
        add_card(slide, "Sensors", lookup["Sensors"], 8.46, 3.12, 4.17, 1.54, accent=BLUE, body_size=8.5)
        add_card(slide, "Performance Parameters", lookup["Performance Parameters"], 8.46, 4.84, 4.17, 1.80, accent=BLUE, body_size=8.5)
    elif page in {17, 18, 20, 25, 26, 27, 28, 29}:
        parameter = next((section for section in sections if section[0] == "Parameters"), None)
        support = [section for section in sections if section[0] != "Parameters"]
        support_h = 5.42 / max(1, len(support))
        for i, (heading, content) in enumerate(support):
            add_card(slide, heading, content, content_x, 1.40 + i * support_h, 4.12, support_h - 0.18, body_size=9.0 if len(content) < 10 else 8.2)
        if parameter:
            param_size = 7.4 if len(parameter[1]) > 25 else 8.0 if len(parameter[1]) > 18 else 8.7
            add_card(slide, parameter[0], parameter[1], 8.46, 1.40, 4.17, 5.24, accent=BLUE, body_size=param_size)
    elif count <= 2:
        heights = [2.52, 2.52]
        positions = [(content_x, 1.40), (content_x, 4.10)]
        for i, (heading, content) in enumerate(sections):
            add_card(slide, heading or name, content, positions[i][0], positions[i][1], content_w, heights[i], body_size=10.2)
    else:
        cols = 2
        rows = math.ceil(count / cols)
        card_h = 5.42 / rows
        for i, (heading, content) in enumerate(sections):
            col, row = i % cols, i // cols
            x = content_x + col * 4.38; y = 1.40 + row * card_h
            add_card(slide, heading or name, content, x, y, 4.14, card_h - 0.18, accent=CYAN if col == 0 else BLUE, body_size=8.7 if len(content) > 8 else 9.5)
    return slide


def build_greeting_star(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 19, texts, section="GREETINGBOT")
    add_title(slide, "GreetingBot Star Key Specifications", compact=True)
    body = clean_footer_items(text_without(texts, ["GreetingBot Star Key Specifications"]))
    rect(slide, 0.58, 1.28, 3.25, 5.70, PANEL, GRID)
    add_picture_contain(slide, assets[0], 0.85, 1.52, 2.72, 3.78)
    add_picture_contain(slide, assets[3], 0.78, 5.45, 1.40, 0.82)
    add_picture_contain(slide, assets[4], 2.22, 5.45, 1.40, 0.82)
    rows = pairs(body)
    for i, (label, value) in enumerate(rows):
        y = 1.35 + i * 0.44
        rect(slide, 4.08, y, 8.50, 0.35, PANEL if i % 2 == 0 else BG_2, GRID)
        add_text(slide, label, 4.22, y + 0.06, 2.12, 0.20, size=8.6, color=CYAN, bold=True)
        add_text(slide, value, 6.45, y + 0.06, 5.95, 0.20, size=8.6, color=WHITE if i < 4 else TEXT, bold=i < 4)
    return slide


def build_humanoid_platform(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 21, texts, section="HUMANOID PLATFORM")
    add_title(slide, "FMC³ Brain+ODM", eyebrow="Products", subtitle="Full-body embodied intelligence platform")
    body = clean_footer_items(text_without(texts, ["Products", "FMC³ Brain+ODM"]))
    heads = ["Seven-degree-of-freedom bionic arm", "Omnidirectional Mobile Chassis", "Power management system", "Embodied Intelligence", "Quick-change End Effector", "Full-body Motion Control System", "Navigation system"]
    sections = split_sections(body, heads)
    add_picture_contain(slide, assets[1], 5.22, 1.45, 2.85, 4.95)
    left_sections, right_sections = sections[:3], sections[3:]
    for i, (head, content) in enumerate(left_sections):
        add_card(slide, head, content, 0.62, 1.46 + i * 1.72, 4.10, 1.48, body_size=9.0)
    for i, (head, content) in enumerate(right_sections):
        add_card(slide, head, content, 8.53, 1.46 + i * 1.30, 4.10, 1.10, accent=BLUE, body_size=8.5, title_size=9.2)
    return slide


def build_scenarios(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 22, texts, section="APPLICATIONS")
    add_title(slide, "Application Scenario")
    source_labels = clean_footer_items(text_without(texts, ["Application Scenario"]))
    labels = [source_labels[i] for i in [0, 2, 4, 5, 3, 1]]
    images = assets[1:7]
    for i, (label, image) in enumerate(zip(labels, images)):
        col, row = i % 3, i // 3
        x, y = 0.62 + col * 4.15, 1.38 + row * 2.65
        rect(slide, x, y, 3.75, 2.34, PANEL, GRID)
        add_picture_contain(slide, image, x + 0.10, y + 0.10, 3.55, 1.72)
        add_text(slide, label, x + 0.12, y + 1.91, 3.51, 0.30, size=10.2, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    return slide


def build_upper_limb(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 23, texts, section="COMPLIANT CONTROL")
    title = "Humanoid upper limb, compliant control, safe human-robot collaboration"
    add_title(slide, title, compact=True)
    body = clean_footer_items(text_without(texts, [title]))
    add_picture_contain(slide, assets[3], 0.62, 1.35, 4.5, 2.25)
    add_metric(slide, "3.5 kg / 3.5 kg", "Rated / Max. load", 0.70, 3.72, 2.15)
    add_metric(slide, "7.5 kg", "Weight", 2.98, 3.72, 2.15, accent=BLUE)
    image_cards = [(assets[0], body[4]), (assets[1], body[5]), (assets[2], body[3])]
    for i, (image, description) in enumerate(image_cards):
        x = 0.68 + i * 1.48
        add_picture_contain(slide, image, x, 4.72, 1.34, 0.90)
        add_text(slide, description, x, 5.66, 1.34, 0.78, size=7.3, color=TEXT)
    features = body[:3] + body[6:]
    for i, item in enumerate(features):
        y = 1.38 + i * 0.83
        rect(slide, 5.48, y, 7.12, 0.67, PANEL if i % 2 == 0 else BG_2, GRID)
        add_text(slide, f"0{i+1}", 5.65, y + 0.17, 0.42, 0.22, size=8.2, color=CYAN, bold=True, font=MONO)
        add_text(slide, item, 6.18, y + 0.10, 6.20, 0.44, size=8.7, color=WHITE if i < 3 else TEXT, bold=i < 3)
    return slide


def build_architecture(prs, texts, assets):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_background(slide, 24, texts, section="SYSTEM ARCHITECTURE")
    title = "Refined and User-Friendly Infrastructure – Customer focus is on AI model deployment and research"
    add_title(slide, title, compact=True)
    body = clean_footer_items(text_without(texts, [title]))
    # Core flow
    nodes = ["Vision System Data", "High Computational Power – NVIDIA Thor", "Embodied AI Model", "Interface Invocation", "Motion Control System", "Full-body Motion Control"]
    positions = [(0.65, 1.55), (3.05, 1.55), (5.45, 1.55), (2.0, 3.0), (4.4, 3.0), (6.8, 3.0)]
    for i, (node, (x, y)) in enumerate(zip(nodes, positions)):
        rect(slide, x, y, 2.05, 0.72, PANEL_2 if i in (1, 2) else PANEL, CYAN if i < 3 else BLUE)
        add_text(slide, node, x + 0.10, y + 0.14, 1.85, 0.38, size=8.6, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for a, b in [(positions[0], positions[1]), (positions[1], positions[2]), (positions[1], positions[3]), (positions[3], positions[4]), (positions[4], positions[5])]:
        line(slide, a[0] + 2.05, a[1] + 0.36, b[0], b[1] + 0.36, CYAN, 1.0)
    add_picture_contain(slide, assets[2], 3.58, 2.34, 1.0, 0.60)
    add_picture_contain(slide, assets[3], 0.70, 2.42, 1.08, 0.60)
    add_picture_contain(slide, assets[4], 7.35, 2.35, 1.15, 0.52)
    add_picture_contain(slide, assets[1], 2.78, 3.84, 0.86, 0.72)
    add_text(slide, ["Ethernet interface", "（Motion Control Message Interface)"], 2.02, 3.78, 2.25, 0.58, size=7.7, color=TEXT, align=PP_ALIGN.CENTER)
    # Robot branches
    branch_nodes = ["Seven-Axis Bionic Arm", "Folding Mechanism", "Omni-directional Chassis", "Motion Optimization and Coordinated Movement"]
    for i, node in enumerate(branch_nodes):
        y = 1.55 + i * 1.05
        rect(slide, 9.45, y, 3.05, 0.72, PANEL, GRID)
        add_text(slide, node, 9.60, y + 0.15, 2.75, 0.38, size=8.7, color=WHITE, bold=True)
        line(slide, 8.85, 3.36, 9.45, y + 0.36, BLUE, 0.8)
    detail = [item for item in body if item not in nodes + branch_nodes and item not in {"Ethernet interface", "（Motion Control Message Interface)"}]
    add_text(slide, detail, 0.72, 5.33, 11.70, 1.28, size=8.2, color=TEXT, bullet=True)
    return slide


def build_end(prs, texts, assets, runtime):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    hero = crop_cover(assets[0], runtime / "end-hero.jpg")
    slide.shapes.add_picture(str(hero), 0, 0, width=Inches(W), height=Inches(H))
    rect(slide, 0, 0, 7.8, H, BG, None, transparency=18)
    add_text(slide, "FMC³ Robotics", 0.72, 1.48, 6.25, 0.72, size=27, color=CYAN, bold=True)
    add_text(slide, "Thanks", 0.68, 2.32, 6.35, 1.18, size=56, color=WHITE, bold=True)
    add_text(slide, "www.fmc3robotics.ai", 0.75, 4.13, 3.2, 0.32, size=12, color=TEXT, font=MONO)
    add_logo(slide)
    add_text(slide, "30 / 30", 11.55, 7.18, 1.15, 0.2, size=8.5, color=CYAN, bold=True, align=PP_ALIGN.RIGHT, font=MONO)
    return slide


def build_deck(output_dir: Path) -> tuple[Path, list[dict[str, object]], dict[int, list[Path]]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    runtime = output_dir / "runtime"; runtime.mkdir(exist_ok=True)
    source_prs = Presentation(SOURCE)
    snapshot = workflow.source_snapshot(source_prs)
    source_text = {int(page["page"]): list(page["text"]) for page in snapshot}
    source_notes = {int(page["page"]): list(page["notes"]) for page in snapshot}
    asset_map = extract_assets(output_dir)
    assets = {page: raster_assets(asset_map, page) for page in range(1, 31)}

    prs = Presentation()
    prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    # Remove the default slide if a template created one.
    while prs.slides:
        slide_id = prs.slides._sldIdLst[0]
        prs.part.drop_rel(slide_id.rId)
        del prs.slides._sldIdLst[0]

    builders = {
        1: lambda: build_cover(prs, source_text[1], assets[1], runtime),
        2: lambda: build_positioning(prs, source_text[2], assets[2], runtime),
        3: lambda: build_brain(prs, source_text[3], assets[3]),
        4: lambda: build_team(prs, source_text[4], assets[4]),
        5: lambda: build_gdpr(prs, source_text[5], assets[5]),
        6: lambda: build_safety(prs, source_text[6], assets[6]),
        7: lambda: build_partners(prs, source_text[7], assets[7]),
        8: lambda: build_portfolio(prs, source_text[8], assets[8]),
        9: lambda: build_brain_suite(prs, source_text[9], assets[9]),
        10: lambda: build_omnigrab(prs, source_text[10], assets[10]),
        19: lambda: build_greeting_star(prs, source_text[19], assets[19]),
        21: lambda: build_humanoid_platform(prs, source_text[21], assets[21]),
        22: lambda: build_scenarios(prs, source_text[22], assets[22]),
        23: lambda: build_upper_limb(prs, source_text[23], assets[23]),
        24: lambda: build_architecture(prs, source_text[24], assets[24]),
        30: lambda: build_end(prs, source_text[30], assets[30], runtime),
    }
    cleaning_asset = assets[8][3]
    for page in range(1, 31):
        if page in builders:
            slide = builders[page]()
        else:
            slide = build_product_page(prs, page, source_text[page], assets[page], cleaning_asset)
        set_notes(slide, source_notes[page])

    pptx_path = output_dir / f"{BASENAME}.pptx"
    prs.save(pptx_path)
    image_map = {page: [path.name for path in asset_map[page]] for page in asset_map}
    workflow.write_workflow_files(output_dir, snapshot, {str(k): v for k, v in image_map.items()})
    return pptx_path, snapshot, asset_map


def render(output_dir: Path, pptx_path: Path) -> tuple[Path, list[Path]]:
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(pptx_path)],
        check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
    )
    pdf = output_dir / f"{BASENAME}.pdf"
    png_dir = output_dir / "png"; png_dir.mkdir(exist_ok=True)
    for stale in png_dir.glob("slide-*.png"):
        stale.unlink()
    subprocess.run(["pdftoppm", "-png", "-r", "144", str(pdf), str(png_dir / "slide")], check=True)
    return pdf, sorted(png_dir.glob("slide-*.png"))


def output_snapshot(path: Path) -> list[dict[str, object]]:
    return workflow.source_snapshot(Presentation(path))


def semantic_values(values: list[str]) -> set[str]:
    return {
        normalize(value) for value in values
        if normalize(value)
        and not re.fullmatch(r"\d+", normalize(value))
        and normalize(value) not in {"The Embodied AI Brain from Europe", "Thanks"}
    }


def audit_content(source_snapshot, output_path: Path, pngs: list[Path], asset_map) -> dict[str, object]:
    target = output_snapshot(output_path)
    pages = []
    for source_page, output_page in zip(source_snapshot, target):
        source_values = semantic_values(list(source_page["text"]))
        output_values = semantic_values(list(output_page["text"]))
        missing = sorted(source_values - output_values)
        pages.append({"page": source_page["page"], "missing_text": missing, "ok": not missing})
    png_checks = {}
    for index, path in enumerate(pngs, 1):
        with Image.open(path) as image:
            png_checks[str(index)] = {"dimensions": list(image.size), "bytes": path.stat().st_size}
    ok = len(target) == 30 and len(pngs) == 30 and all(page["ok"] for page in pages)
    return {
        "ok": ok,
        "source": str(SOURCE),
        "output": str(output_path),
        "slide_count": len(target),
        "page_checks": pages,
        "source_slide_asset_relations": sum(len(paths) for paths in asset_map.values()),
        "png_checks": png_checks,
    }


def build_preview(output_dir: Path, pngs: list[Path]) -> None:
    figures = "".join(
        f'<figure><img src="png/{html.escape(path.name)}" alt="Slide {i}"><figcaption>{i:02d}</figcaption></figure>'
        for i, path in enumerate(pngs, 1)
    )
    write_text(output_dir / "preview.html", f"""<!doctype html><html><head><meta charset="utf-8"><title>FMC3 Product Manual · Dark Tech Redesign</title><style>*{{box-sizing:border-box}}body{{margin:0;background:#020611;color:white;font-family:Arial,sans-serif}}header{{position:sticky;top:0;z-index:2;padding:16px 26px;background:#050b1f;border-bottom:1px solid #16465c}}main{{padding:24px;display:grid;grid-template-columns:repeat(auto-fit,minmax(420px,1fr));gap:22px}}figure{{margin:0}}img{{width:100%;display:block;border:1px solid #16465c}}figcaption{{padding-top:7px;color:#22d3ee;font-family:monospace}}</style></head><body><header><strong>FMC³ Robotics · Product Manual · Dark Tech Redesign</strong></header><main>{figures}</main></body></html>""")


def write_manifest(output_dir: Path) -> None:
    write_json(output_dir / "delivery-manifest.json", {
        "run_id": output_dir.name,
        "generated_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "summary": {"total_pages": 30, "format": "16:9 widescreen", "build": "blank-canvas redesign"},
        "artifacts": {
            "preview_html": "preview.html",
            "presentation_pptx": f"{BASENAME}.pptx",
            "presentation_png_pptx": f"{BASENAME}.pptx",
            "presentation_svg_pptx": f"{BASENAME}.pptx",
            "brochure_pdf": f"{BASENAME}.pdf",
            "content_audit": "content-audit.json",
        },
        "pages": [{"page": page, "planning": f"planning/planning{page}.json", "png": f"png/slide-{page:02d}.png"} for page in range(1, 31)],
        "export_notes": {
            "presentation_pptx": "Editable native PowerPoint rebuilt from a blank canvas.",
            "content_policy": "Source wording, data, order, notes, and meaningful supplied visuals preserved.",
        },
    })


def build(output_dir: Path) -> None:
    pptx_path, snapshot, asset_map = build_deck(output_dir)
    pdf_path, pngs = render(output_dir, pptx_path)
    audit = audit_content(snapshot, pptx_path, pngs, asset_map)
    write_json(output_dir / "content-audit.json", audit)
    build_preview(output_dir, pngs)
    write_manifest(output_dir)
    if not audit["ok"]:
        raise RuntimeError(f"Content audit failed: {output_dir / 'content-audit.json'}")
    print(f"Built: {pptx_path}")
    print(f"Built: {pdf_path}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Rebuild the FMC3 product manual in dark-tech style")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.output_dir.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
