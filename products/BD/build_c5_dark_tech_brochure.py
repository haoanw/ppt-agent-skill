#!/usr/bin/env python3
"""Build the FMC3 C5 German business brochure in the dark_tech style.

The source DOCX remains untouched. The generated package contains:
- workflow contracts and per-page planning JSON
- four A4 portrait HTML pages
- high-resolution PNG pages
- A4 PDF and PPTX exports
- a content audit against the DOCX text and embedded media
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "products/BD/FMC3-C5_Deutsch_Broschure_Business.docx"
DEFAULT_OUTPUT = ROOT / "products/BD/FMC3-C5_Deutsch_Broschure_Business_dark_tech"
CHROME = shutil.which("google-chrome") or shutil.which("chromium")

PAGE_W = 1120
PAGE_H = 1584
TOTAL_PAGES = 4

WORKFLOW_METADATA = {
    "stage": "planning",
    "workflow_version": "2026.04.09-v4.1",
    "planning_schema_version": "4.1",
    "planning_packet_version": "4.1",
    "planning_continuity_version": "4.1",
}

FEATURES = [
    "Doppelwalzenbürsten – Kehren und Schrubben in einem Arbeitsgang",
    "Zwei-Kammer-Saugfuß für gründliche Wasseraufnahme, weniger Restwasser in Fugen",
    "Laser- und visuelle Lokalisierung für flexible Einsatzumgebungen",
    "Autonomes Laden sowie Frisch- und Schmutzwasser-Management",
    "Tür- und Aufzugsanbindung für den gebäudeweiten Einsatz",
    "Selbstreinigender Schmutzwassertank – täglicher Wartungsaufwand halbiert",
    "Randbürste reinigt dicht an Kanten und Rändern",
    "Hochwertige Komponenten für niedrige Lebenszykluskosten",
]

PAGE_TEXT = {
    1: [
        "Einladung zur Demonstration des C5-Reinigungsroboters von FMC3",
        "Intelligenter gewerblicher Reinigungsroboter für mittlere und große Flächen",
        "Sehr geehrte Damen und Herren,",
        "wir freuen uns, Ihnen unseren Reinigungsroboter C5 vorstellen zu dürfen, und laden Sie herzlich ein, ihn persönlich kennenzulernen.",
        "Der C5 ist eine Revolution in der Reinigungstechnik: Mit ihm bieten Sie Ihren Kunden eine präzise und effiziente Reinigung verschiedenster Innenflächen – und damit Dienstleistungen auf höchstem Niveau.",
        "Der C5 auf einen Blick",
        *FEATURES,
        "1.980 m²/h",
        "theoretische Flächenleistung",
        "550 mm",
        "Schrubbbreite",
        "< 68 dB(A)",
        "Geräuschpegel",
        "90 L",
        "großer Wassertank",
    ],
    2: [
        "Wirtschaftlichkeit auf einen Blick",
        "Präzision und Effizienz stehen bei uns an erster Stelle. Mit dem C5 sparen Sie Zeit und Arbeitskosten und senken zugleich den Verbrauch von Wasser und Reinigungsmitteln – eine Beispielrechnung gegenüber der manuellen Reinigung:",
        "9.089 €",
        "Einsparungen pro Jahr",
        "1,3 Jahre",
        "Amortisationszeit (ROI)",
        "1,54 Mio. m²",
        "gereinigte Fläche pro Jahr",
        "Beispielrechnung: jährlicher Kostenvergleich – manuelle Reinigung vs. C5 (60 Monate, 1.000 h/Jahr)",
        "Über FMC3 Robotics",
        "FMC3 steht für die Zukunft der Robotik. Von unserem Standort in Ingolstadt aus bieten wir Ihnen die neuesten und innovativsten Technologien an. Unser Unternehmen wurde im Oktober 2025 von Herrn Prof. Dr. Bjoern Giesler, Herrn Dr. Wang Cheng und Herrn Dr. David Fan gemeinsam mit dem Ziel gegründet, Europa mit moderner Robotiktechnologie zu versorgen.",
        "Seitdem haben wir bereits zahlreiche Verträge mit Kunden in Deutschland und Europa abgeschlossen. Darüber hinaus versorgen wir unsere Kunden mit Begrüßungs- und Lieferrobotern. Unsere Ingenieure in Shanghai arbeiten derzeit an einem humanoiden Roboter; außerdem starten wir eine Zusammenarbeit im Bereich der Sicherheitsroboterhunde.",
        "FMC3 auf einen Blick",
        "Gegründet: Oktober 2025",
        "Hauptsitz: Ingolstadt, Deutschland",
        "Entwicklung: Ingolstadt & Shanghai",
        "Lieferkapazität: bis zu 30 Roboter/Monat",
        "Portfolio: Reinigungs-, Begrüßungs- und Lieferroboter",
    ],
    3: [
        "Unser Kernteam",
        "Hinter FMC3 steht ein Team europäischer Wissenschaftler für intelligente Systeme – mit langjähriger Erfahrung in der Automobilindustrie, der KI-Forschung und dem internationalen B2B-Geschäft.",
        "Das Kernteam von FMC3 Robotics",
        "Produktparameter",
        "Kompakte Abmessungen, 170 kg Gesamtgewicht und eine Arbeitsstation mit automatischem Frisch- und Schmutzwasser-Management – alle technischen Daten im Überblick:",
        "Technische Daten des C5 und der Arbeitsstation",
    ],
    4: [
        "Einsatzszenarien, Service und Support",
        "Anwendungsszenarien, geeignete Bodenbeläge sowie Service- und Supportsystem",
        "Lieferung & Inbetriebnahme",
        "Wir sind in der Lage, bis zu 30 Roboter pro Monat auszuliefern, und kümmern uns um den gesamten Prozess – von den Tests über die Installation bis zur Inbetriebnahme.",
        "Service & After-Sales",
        "Unsere Ingenieure in Ingolstadt stehen Ihnen täglich zur Verfügung – inklusive Schulungen, kurzer Reaktionszeiten und kostenloser Erstwartung.",
        "Erleben Sie den C5 live",
        "Wenn Sie Interesse an unserem C5 haben, senden wir Ihnen gerne weitere Informationen zu und laden Sie herzlich zu einer Demonstration des Roboters ein – bei Ihnen vor Ort oder in unserem Hauptsitz in Ingolstadt.",
        "Wir freuen uns auf Ihre Rückmeldung.",
        "Mit freundlichen Grüßen",
        "Haoan Wang",
        "Ihr Ansprechpartner bei FMC3 Robotics",
    ],
}

FOOTER_TEXT = "FMC3 Robotics · Ingolstadt, Deutschland · www.FMC3-robotics.ai"


class TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


def normalize_text(value: str) -> str:
    value = html.unescape(value).replace("\u00a0", " ")
    value = re.sub(r"\s+", " ", value).strip()
    return re.sub(r"^[▪•]\s*", "", value)


def iter_doc_blocks(document: DocxDocument):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def source_text_segments(source: Path) -> list[str]:
    doc = Document(source)
    segments: list[str] = []
    for block in iter_doc_blocks(doc):
        paragraphs = [block] if isinstance(block, Paragraph) else [
            paragraph
            for row in block.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        ]
        for paragraph in paragraphs:
            text = normalize_text(paragraph.text)
            if text and text not in segments:
                segments.append(text)
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            text = normalize_text(paragraph.text)
            if text and text not in segments:
                segments.append(text)
    return segments


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def write_json(path: Path, payload: object) -> None:
    write_text(path, json.dumps(payload, ensure_ascii=False, indent=2) + "\n")


def extract_media(source: Path, images_dir: Path) -> dict[str, dict[str, str | int]]:
    images_dir.mkdir(parents=True, exist_ok=True)
    inventory: dict[str, dict[str, str | int]] = {}
    with ZipFile(source) as archive:
        for name in sorted(item for item in archive.namelist() if item.startswith("word/media/")):
            filename = Path(name).name
            data = archive.read(name)
            destination = images_dir / filename
            destination.write_bytes(data)
            inventory[filename] = {
                "source_member": name,
                "bytes": len(data),
                "sha256": sha256_bytes(data),
            }
    return inventory


def workflow_files(output_dir: Path) -> None:
    write_text(
        output_dir / "interview-qa.txt",
        """# Interview QA

scenario: Deutsche B2B-Produktbroschüre und Einladung zur C5-Demonstration
audience: Gebäudereiniger, Facility-Management, gewerbliche Entscheider und potenzielle Vertriebspartner
target_action: Interesse wecken und eine persönliche C5-Demonstration vereinbaren
expected_pages: 4
page_density: balanced
style: dark_tech
brand: FMC3 Robotics, vorhandenes Logo und alle Originalinhalte beibehalten
must_include: Sämtliche Texte, Kennzahlen, Bilder, Teamprofile, Produktparameter, Einsatzszenarien, Services und Kontaktdaten aus der DOCX
must_avoid: Keine inhaltlichen Kürzungen, keine erfundenen Aussagen, keine Änderung der deutschen Sprache
language: Deutsch
imagery: provided
material_strategy: Ausschließlich die bereitgestellte DOCX und ihre eingebetteten Medien
subagent_model_strategy: Hauptagent
subagent_thinking_effort: hoch
manual_audit_mode: fine_grained
manual_audit_scope: page_html, page_review
manual_audit_assets: planning, html, png, content-audit
""",
    )
    write_text(
        output_dir / "requirements-interview.txt",
        """# Requirements Interview

scenario: Deutsche B2B-Produktbroschüre und Einladung zur C5-Demonstration
audience: Gebäudereiniger, Facility-Management, gewerbliche Entscheider und potenzielle Vertriebspartner
target_action: Interesse wecken und eine persönliche C5-Demonstration vereinbaren
expected_pages: 4
page_density: balanced
style: dark_tech
brand: FMC3 Robotics; vorhandenes Logo, Bildmaterial und alle Originalinhalte beibehalten
must_include: Vollständiger Inhalt der Quelldatei FMC3-C5_Deutsch_Broschure_Business.docx
must_avoid: Keine Kürzungen, keine Übersetzungen, keine neuen Leistungsversprechen
language: Deutsch
imagery: provided
material_strategy: non-research; nur bereitgestellte Quelle
subagent_model_strategy: Hauptagent
subagent_thinking_effort: hoch
manual_audit_mode: fine_grained
manual_audit_scope: page_html, page_review
manual_audit_assets: planning, html, png, content-audit
density_bias: balanced
branch: 非research
""",
    )
    source_lines = [
        "# Source Brief",
        "",
        "topic: FMC3 C5 Reinigungsroboter, deutsche Business-Broschüre",
        "usage: B2B-Produktvorstellung und Einladung zu einer Live-Demonstration",
        "constraints: Jeder Textbaustein und jedes eingebettete Medium muss erhalten bleiben.",
        "risk: Die Informationsgrafiken enthalten zusätzlichen Text; deshalb werden sie byte-identisch übernommen.",
        "",
        "Thema: FMC3 C5 Reinigungsroboter, deutsche Business-Broschüre",
        "Quelle: products/BD/FMC3-C5_Deutsch_Broschure_Business.docx",
        "Einsatz: B2B-Produktvorstellung und Einladung zu einer Live-Demonstration",
        "Einschränkung: Jeder Textbaustein und jedes eingebettete Medium muss erhalten bleiben.",
        "Risiko: Die Informationsgrafiken enthalten zusätzlichen Text; deshalb werden sie byte-identisch übernommen.",
        "",
    ]
    for page, segments in PAGE_TEXT.items():
        source_lines.extend([f"## Seite {page}", ""])
        source_lines.extend(f"- {segment}" for segment in segments)
        source_lines.append("")
    source_lines.extend(["## Wiederkehrende Fußzeile", "", f"- {FOOTER_TEXT}", ""])
    write_text(output_dir / "source-brief.txt", "\n".join(source_lines))

    write_text(
        output_dir / "outline.txt",
        """# Outline

密度倾向: balanced
密度曲线: mid_low -> medium -> medium -> medium

### 第 1 页: C5 Produktauftakt
页目标: Die Einladung und den C5 als präzise gewerbliche Reinigungslösung positionieren
页面类型映射: cover
密度下限: low
密度目标: mid_low
密度上限: medium
节奏动作: 铺垫
信息姿态: 结论页
锚点类型: 图片

### 第 2 页: Wirtschaftlichkeit und FMC3
页目标: Wirtschaftlichen Nutzen belegen und FMC3 als europäischen Robotikpartner einordnen
页面类型映射: content
密度下限: mid_low
密度目标: medium
密度上限: high
节奏动作: 爆发
信息姿态: 证据页
锚点类型: KPI

### 第 3 页: Team und Produktparameter
页目标: Kompetenz des Kernteams und technische Leistungsfähigkeit des C5 zeigen
页面类型映射: content
密度下限: mid_low
密度目标: medium
密度上限: high
节奏动作: 推进
信息姿态: 解释页
锚点类型: 图片

### 第 4 页: Einsatz und Einladung
页目标: Einsatzbreite und Serviceversprechen in eine klare Demonstrationseinladung überführen
页面类型映射: content
密度下限: mid_low
密度目标: medium
密度上限: high
节奏动作: 收束
信息姿态: 结论页
锚点类型: 图片

自审通过: Alle vier Originalseiten und ihre Inhalte sind vollständig abgebildet.
""",
    )

    style = {
        "style_id": "dark_tech",
        "style_name": "暗黑科技 (Dark Tech)",
        "mood_keywords": ["深空冷寂", "精密仪器", "微光脉搏", "数据洪流", "未来感"],
        "design_soul": "深蓝黑幕中的 C5 像一台被冷青扫描光逐层唤醒的精密仪器，数据、产品与团队证据以清晰的工业节奏展开。",
        "variation_strategy": "封面以产品大图和深空留白聚焦；经济页用青色 KPI 与证据面板形成高密扫描；团队和参数页采用上下叙事；末页以场景图、服务模块和发光 CTA 收束。背景、字体、青色边线和页脚锁定不变。",
        "decoration_dna": {
            "signature_move": "深空网格 + 冷青扫描线 + L 形角标 + 半透明数据面板",
            "forbidden": ["花卉装饰", "马卡龙色", "厚重圆角胶囊", "波浪分隔线", "无关装饰图片"],
            "recommended_combos": [
                "网格底纹 + 青色角标 + 产品大图",
                "微光边框 + 超级 KPI + 等宽标签",
                "低透明辉光 + 信息图原图 + 细线页脚",
            ],
        },
        "font_family": "'Noto Sans', 'Liberation Sans', Arial, sans-serif",
        "css_variables": {
            "bg_primary": "#050b1f",
            "bg_secondary": "#0a1f3d",
            "card_bg_from": "rgba(34,211,238,0.10)",
            "card_bg_to": "rgba(99,102,241,0.06)",
            "card_border": "rgba(34,211,238,0.24)",
            "card_radius": "8px",
            "text_primary": "#FFFFFF",
            "text_secondary": "rgba(255,255,255,0.68)",
            "accent_1": "#22D3EE",
            "accent_2": "#3B82F6",
            "accent_3": "#6366F1",
            "accent_4": "#FDE047",
        },
        "css_snippets": {
            "panel": "background:linear-gradient(145deg,rgba(34,211,238,.10),rgba(99,102,241,.05));border:1px solid rgba(34,211,238,.24);border-radius:8px;",
            "label": "font-family:'Noto Sans Mono',monospace;letter-spacing:.16em;text-transform:uppercase;color:#22D3EE;",
        },
    }
    write_json(output_dir / "style.json", style)


def image_contract(filename: str | None = None, placement: str = "inline") -> dict:
    if filename is None:
        return {
            "mode": "decorate",
            "needed": False,
            "usage": None,
            "placement": None,
            "content_description": None,
            "source_hint": None,
            "decorate_brief": "Dark-tech grid and cyan scan-line treatment only.",
        }
    return {
        "mode": "provided",
        "needed": True,
        "usage": "inline-illustration",
        "placement": placement,
        "content_description": "Original FMC3 brochure visual, preserved without content changes.",
        "source_hint": f"images/{filename}",
        "decorate_brief": "Place inside a restrained cyan technical frame; do not crop away text.",
    }


def card(
    card_id: str,
    role: str,
    card_type: str,
    card_style: str,
    argument_role: str,
    headline: str,
    body: list[str],
    image: dict,
    body_max_lines: int,
) -> dict:
    return {
        "card_id": card_id,
        "role": role,
        "card_type": card_type,
        "card_style": card_style,
        "argument_role": argument_role,
        "headline": headline,
        "body": body,
        "data_points": [],
        "content_budget": {
            "headline_max_chars": 42,
            "body_max_bullets": max(1, min(8, len(body))),
            "body_max_lines": body_max_lines,
        },
        "image": image,
        "resource_ref": {"principle": "visual-hierarchy" if role == "anchor" else "composition"},
    }


def planning_payloads(output_dir: Path) -> list[dict]:
    common = {
        "density_contract": {
            "deck_bias": "balanced",
            "page_lower_bound": "mid_low",
            "page_upper_bound": "high",
            "max_cards": 4,
            "max_charts": 2,
            "min_body_font_px": 18,
            "max_lines_per_card": 5,
            "image_policy": "support_only",
            "decoration_budget": "medium",
            "overflow_strategy": "tighten_budget",
        },
        "density_label": "medium",
        "negative_space_target": "medium",
        "page_text_strategy": "Deutsche Originaltexte unverändert, klare technische Hierarchie und kurze scanbare Module.",
        "must_avoid": ["Textkürzung", "Informationsgrafiken beschneiden", "zu geringe Schriftgröße", "dekorative Bilder ohne Quellenbezug"],
        "variation_guardrails": {
            "same_gene_as_deck": "Dunkelblauer Grund, Cyan-Linien, weiße Typografie, technische Footer-Leiste.",
            "different_from_previous": [],
        },
        "decoration_hints": {
            "background": {"feel": "tiefer technischer Raum", "restraint": "sehr niedrige Deckkraft", "techniques": ["grid", "scan-line"]},
            "floating": {"feel": "präzise Instrumentenmarker", "restraint": "maximal zwei Akzente", "techniques": ["corner-lines"]},
            "page_accent": {"feel": "kaltes Cyan", "restraint": "nur Fokusdaten", "techniques": ["cyan-rule"]},
        },
        "source_guidance": {
            "brief_sections": ["source-brief.txt"],
            "citation_expectation": "Nur Inhalte und Originalmedien aus der bereitgestellten DOCX.",
            "strictness": "Keine Ergänzungen oder Umformulierungen.",
        },
    }

    page1 = {
        "page": {
            "slide_number": 1,
            "page_type": "cover",
            "narrative_role": "cover",
            "title": PAGE_TEXT[1][0],
            "page_goal": "Den C5 als präzise gewerbliche Reinigungslösung vorstellen und zur Demonstration einladen.",
            "audience_takeaway": "Der C5 verbindet autonome Reinigung mit hoher Flächenleistung und niedrigem Wartungsaufwand.",
            "visual_weight": 7,
            "density_label": "mid_low",
            "density_reason": "Die Titelseite bewahrt Einleitung, acht Produktvorteile, vier Kennzahlen und das Originalproduktbild, bleibt aber durch eine klare Zweiteilung scanbar.",
            "density_contract": {
                "deck_bias": "balanced",
                "page_lower_bound": "low",
                "page_upper_bound": "medium",
                "max_cards": 3,
                "max_charts": 1,
                "min_body_font_px": 20,
                "max_lines_per_card": 4,
                "image_policy": "flexible",
                "decoration_budget": "medium",
                "overflow_strategy": "rebalance_layout",
            },
            "focus_zone": "oberes rechtes Drittel auf dem C5, danach Titel und Leistungskennzahlen",
            "negative_space_target": "medium",
            "page_text_strategy": "Originaleinleitung als kurze Absätze, Produktvorteile zweispaltig, Kennzahlen als Instrumentenleiste.",
            "rhythm_action": "铺垫",
            "must_avoid": ["Titel in eine Karte sperren", "Produktbild klein darstellen", "Produktvorteile kürzen"],
            "variation_guardrails": {
                "same_gene_as_deck": "Dark-tech Hintergrund, Cyan-Linien, weißes Logo, technische Fußzeile.",
                "different_from_previous": ["Erste Seite nutzt eine offene Hero-Komposition ohne Standardheader."],
            },
            "director_command": {
                "mood": "präziser Produktauftakt",
                "spatial_strategy": "Titel links, Produkt rechts, Funktionen und Kennzahlen als untere technische Ebene",
                "anchor_treatment": "C5 als freigestelltes Hero-Objekt in cyan gerahmtem Sichtfenster",
                "techniques": ["grid", "corner-lines", "data-rail"],
                "prose": "Der C5 tritt aus einer dunklen Testkammer in einen kühlen Cyan-Scan.",
            },
            "decoration_hints": common["decoration_hints"],
            "resources": {
                "page_template": "cover",
                "layout_refs": [],
                "block_refs": [],
                "chart_refs": [],
                "principle_refs": ["color-psychology", "visual-hierarchy", "composition"],
                "resource_rationale": "Das freie Cover trägt Produktbild, Einladung und Kennzahlen ohne eine starre Inhaltsseitenstruktur.",
            },
            "cards": [
                card("s01-anchor-1", "anchor", "image_hero", "accent", "claim", PAGE_TEXT[1][0], PAGE_TEXT[1][1:5], image_contract("image1.png", "right-half"), 4),
                card("s01-support-1", "support", "list", "outline", "evidence", "Der C5 auf einen Blick", FEATURES, image_contract(), 4),
                card("s01-context-1", "context", "data_highlight", "glass", "evidence", "Leistungsdaten", PAGE_TEXT[1][14:], image_contract(), 4),
            ],
            "workflow_metadata": WORKFLOW_METADATA,
        }
    }

    page2 = {
        "page": {
            **common,
            "slide_number": 2,
            "page_type": "content",
            "narrative_role": "evidence",
            "title": PAGE_TEXT[2][0],
            "page_goal": "Wirtschaftlichen Nutzen belegen und FMC3 als europäischen Robotikpartner einordnen.",
            "audience_takeaway": "Der C5 amortisiert sich in 1,3 Jahren und wird von einem europäischen Robotikunternehmen betreut.",
            "visual_weight": 8,
            "density_reason": "ROI-Grafik, drei Kennzahlen und zwei Unternehmensblöcke benötigen mittlere Dichte und eine vertikale Beweisführung.",
            "layout_hint": "mixed-grid",
            "layout_variation_note": "Nach dem offenen Produktcover folgt eine datengetriebene Instrumententafel mit klar getrenntem Unternehmensnachweis.",
            "focus_zone": "obere KPI-Leiste und zentrale ROI-Grafik",
            "rhythm_action": "爆发",
            "variation_guardrails": {
                **common["variation_guardrails"],
                "different_from_previous": ["Stärkere Datenrasterung", "ROI-Grafik als zentraler Beweis", "Textbeleg im unteren Drittel"],
            },
            "director_command": {
                "mood": "wirtschaftlicher Beweis",
                "spatial_strategy": "KPI-Leiste, große Beweisgrafik, zweispaltiger Unternehmensblock",
                "anchor_treatment": "9.089 € als leuchtender Primärwert",
                "techniques": ["data-rail", "evidence-frame", "split-proof"],
                "prose": "Eine dunkle Kontrolltafel zeigt zuerst die Einsparung, dann die vollständige Berechnung und schließlich den Anbieter dahinter.",
            },
            "resources": {
                "page_template": None,
                "layout_refs": ["mixed-grid"],
                "block_refs": [],
                "chart_refs": [],
                "principle_refs": ["data-visualization", "visual-hierarchy", "composition"],
                "resource_rationale": "Mixed-grid verbindet drei KPIs, die unveränderte ROI-Grafik und den Unternehmensnachweis auf einer A4-Seite.",
            },
            "cards": [
                card("s02-anchor-1", "anchor", "data_highlight", "accent", "claim", "Wirtschaftlichkeit", PAGE_TEXT[2][1:8], image_contract(), 5),
                card("s02-support-1", "support", "diagram", "outline", "evidence", PAGE_TEXT[2][8], [PAGE_TEXT[2][8]], image_contract("image2.jpeg"), 3),
                card("s02-support-2", "support", "text", "filled", "evidence", PAGE_TEXT[2][9], PAGE_TEXT[2][10:12], image_contract(), 5),
                card("s02-context-1", "context", "list", "glass", "framework", PAGE_TEXT[2][12], PAGE_TEXT[2][13:], image_contract(), 5),
            ],
            "workflow_metadata": WORKFLOW_METADATA,
        }
    }

    page3 = {
        "page": {
            **common,
            "slide_number": 3,
            "page_type": "content",
            "narrative_role": "framework",
            "title": PAGE_TEXT[3][0],
            "page_goal": "Kompetenz des Kernteams und technische Leistungsfähigkeit des C5 zeigen.",
            "audience_takeaway": "FMC3 verbindet europäische Industrie-, KI- und B2B-Erfahrung mit einem klar spezifizierten Reinigungsrobotersystem.",
            "visual_weight": 8,
            "density_reason": "Zwei informationsreiche Originalgrafiken werden vollständig gezeigt und durch exakt erhaltene Einleitungstexte verbunden.",
            "layout_hint": "waterfall",
            "layout_variation_note": "Die Seite wechselt von der KPI-Tafel zu einer vertikalen Beweisfolge aus Team und Produktparametern.",
            "focus_zone": "Teamgrafik im oberen und Produktgrafik im unteren Seitenzentrum",
            "rhythm_action": "推进",
            "variation_guardrails": {
                **common["variation_guardrails"],
                "different_from_previous": ["Vertikale Bildfolge", "weniger Textkarten", "zwei große Quellenvisuals"],
            },
            "director_command": {
                "mood": "Kompetenz und technische Substanz",
                "spatial_strategy": "zwei große Originaltafeln entlang einer vertikalen Scan-Achse",
                "anchor_treatment": "Teamgrafik als Vertrauensanker, Produktparameter als technischer Abschluss",
                "techniques": ["waterfall", "source-frame", "section-rule"],
                "prose": "Der Blick scannt erst die Menschen hinter FMC3 und fällt danach auf die technische Anatomie des C5.",
            },
            "resources": {
                "page_template": None,
                "layout_refs": ["waterfall"],
                "block_refs": [],
                "chart_refs": [],
                "principle_refs": ["visual-hierarchy", "cognitive-load", "composition"],
                "resource_rationale": "Waterfall hält beide textreichen Originalgrafiken groß genug und bewahrt ihre Lesbarkeit.",
            },
            "cards": [
                card("s03-anchor-1", "anchor", "people", "accent", "evidence", PAGE_TEXT[3][0], PAGE_TEXT[3][1:3], image_contract("image3.jpeg"), 4),
                card("s03-support-1", "support", "data", "filled", "evidence", PAGE_TEXT[3][3], PAGE_TEXT[3][4:], image_contract("image4.png"), 4),
                card("s03-context-1", "context", "text", "transparent", "synthesis", "FMC3 Expertise", ["Automobilindustrie", "KI-Forschung", "internationales B2B-Geschäft"], image_contract(), 3),
            ],
            "workflow_metadata": WORKFLOW_METADATA,
        }
    }

    page4 = {
        "page": {
            **common,
            "slide_number": 4,
            "page_type": "content",
            "narrative_role": "cta",
            "title": PAGE_TEXT[4][0],
            "page_goal": "Einsatzbreite und Serviceversprechen in eine klare Demonstrationseinladung überführen.",
            "audience_takeaway": "Der C5 deckt vielfältige gewerbliche Einsatzfelder ab und FMC3 begleitet Lieferung, Schulung und Erstwartung.",
            "visual_weight": 7,
            "density_reason": "Eine große Anwendungsgrafik, zwei Serviceblöcke, CTA und persönliche Signatur werden mit mittlerer Dichte hierarchisch gestaffelt.",
            "layout_hint": "l-shape",
            "layout_variation_note": "Die Schlussseite öffnet sich wieder: großes Szenarienbild oben, Service und CTA als klare Abschlussstufen.",
            "focus_zone": "Anwendungsgrafik und cyan leuchtender Demonstrationsaufruf",
            "rhythm_action": "收束",
            "variation_guardrails": {
                **common["variation_guardrails"],
                "different_from_previous": ["Ein dominantes Szenarienbild", "zwei Servicekarten", "breiter CTA-Abschluss"],
            },
            "director_command": {
                "mood": "verlässlicher Abschluss mit Handlungsimpuls",
                "spatial_strategy": "große Einsatzgrafik, parallele Servicekarten, breite CTA-Leiste, persönliche Signatur",
                "anchor_treatment": "CTA erhält den stärksten Cyan-Rand der gesamten Broschüre",
                "techniques": ["l-shape", "service-pair", "cta-glow"],
                "prose": "Nach der Einsatzübersicht verdichtet sich die Seite zu Serviceversprechen und einer klaren Einladung, den C5 live zu erleben.",
            },
            "resources": {
                "page_template": None,
                "layout_refs": ["l-shape"],
                "block_refs": [],
                "chart_refs": [],
                "principle_refs": ["narrative-arc", "visual-hierarchy", "composition"],
                "resource_rationale": "L-shape lässt die große Szenariengrafik führen und ordnet Service, CTA und Signatur als Abschlussbewegung darunter.",
            },
            "cards": [
                card("s04-anchor-1", "anchor", "diagram", "accent", "evidence", PAGE_TEXT[4][0], PAGE_TEXT[4][1:2], image_contract("image5.png"), 4),
                card("s04-support-1", "support", "process", "outline", "method", PAGE_TEXT[4][2], PAGE_TEXT[4][3:4], image_contract(), 5),
                card("s04-support-2", "support", "text", "filled", "method", PAGE_TEXT[4][4], PAGE_TEXT[4][5:6], image_contract(), 5),
                card("s04-context-1", "context", "quote", "glass", "synthesis", PAGE_TEXT[4][6], PAGE_TEXT[4][7:], image_contract(), 5),
            ],
            "workflow_metadata": WORKFLOW_METADATA,
        }
    }
    return [page1, page2, page3, page4]


def common_css() -> str:
    return f"""
@page {{ size: A4 portrait; margin: 0; }}
:root {{
  --bg-primary:#050b1f;
  --bg-secondary:#0a1f3d;
  --card-bg-from:rgba(34,211,238,.10);
  --card-bg-to:rgba(99,102,241,.055);
  --card-border:rgba(34,211,238,.25);
  --text-primary:#fff;
  --text-secondary:rgba(255,255,255,.70);
  --accent-1:#22D3EE;
  --accent-2:#3B82F6;
  --accent-3:#6366F1;
  --accent-4:#FDE047;
}}
* {{ box-sizing:border-box; }}
html,body {{ width:{PAGE_W}px; height:{PAGE_H}px; margin:0; overflow:hidden; }}
body {{
  color:var(--text-primary);
  font-family:'Noto Sans','Liberation Sans',Arial,sans-serif;
  background:
    radial-gradient(circle at 86% 12%, rgba(99,102,241,.24), transparent 29%),
    radial-gradient(circle at 8% 78%, rgba(34,211,238,.14), transparent 31%),
    linear-gradient(180deg,var(--bg-secondary) 0%,var(--bg-primary) 42%,#030817 100%);
  position:relative;
  -webkit-font-smoothing:antialiased;
}}
.grid {{ position:absolute; inset:0; opacity:.28; background-image:linear-gradient(rgba(255,255,255,.025) 1px,transparent 1px),linear-gradient(90deg,rgba(255,255,255,.025) 1px,transparent 1px); background-size:64px 64px; }}
.scan {{ position:absolute; left:0; right:0; top:182px; height:1px; background:linear-gradient(90deg,transparent,var(--accent-1),transparent); opacity:.7; }}
.corner {{ position:absolute; width:30px; height:30px; opacity:.65; }}
.corner.tl {{ top:24px; left:24px; border-top:1px solid var(--accent-1); border-left:1px solid var(--accent-1); }}
.corner.br {{ right:24px; bottom:24px; border-right:1px solid var(--accent-1); border-bottom:1px solid var(--accent-1); }}
.page {{ position:relative; z-index:2; height:100%; padding:58px 70px 54px; }}
.brand {{ height:62px; display:flex; align-items:center; justify-content:flex-end; margin-bottom:28px; }}
.brand img {{ width:154px; height:auto; filter:invert(1); opacity:.92; }}
.eyebrow {{ display:flex; align-items:center; gap:10px; color:var(--accent-1); font-size:13px; font-weight:700; letter-spacing:.15em; text-transform:uppercase; margin-bottom:12px; }}
.pulse {{ width:7px; height:7px; border-radius:50%; background:var(--accent-1); box-shadow:0 0 14px rgba(34,211,238,.9); }}
h1,h2,h3,p {{ margin:0; }}
h1 {{ font-size:50px; line-height:1.04; font-weight:800; letter-spacing:0; max-width:920px; }}
h2 {{ font-size:28px; line-height:1.18; font-weight:750; letter-spacing:0; }}
h3 {{ font-size:19px; line-height:1.25; font-weight:750; }}
.subhead {{ margin-top:14px; color:var(--text-secondary); font-size:20px; line-height:1.45; max-width:900px; }}
.body {{ color:var(--text-secondary); font-size:16px; line-height:1.62; }}
.body strong {{ color:var(--text-primary); }}
.panel {{ background:linear-gradient(145deg,var(--card-bg-from),var(--card-bg-to)); border:1px solid var(--card-border); border-radius:8px; box-shadow:0 20px 60px rgba(0,0,0,.22); }}
.image-frame {{ background:rgba(255,255,255,.97); border:1px solid rgba(34,211,238,.38); border-radius:8px; padding:8px; box-shadow:0 0 34px rgba(34,211,238,.11); overflow:hidden; }}
.image-frame img {{ width:100%; height:100%; object-fit:contain; display:block; }}
.caption {{ color:rgba(255,255,255,.52); font-size:13px; line-height:1.4; font-style:italic; text-align:center; margin-top:8px; }}
.section-title {{ display:flex; align-items:center; gap:14px; margin-bottom:12px; }}
.section-title .rule {{ flex:1; height:1px; background:linear-gradient(90deg,var(--accent-1),transparent); opacity:.65; }}
.kpi-row {{ display:grid; grid-template-columns:repeat(3,1fr); gap:16px; }}
.kpi {{ padding:21px 18px 18px; min-height:110px; }}
.kpi .value {{ font-size:34px; font-weight:800; line-height:1; color:var(--text-primary); font-variant-numeric:tabular-nums; }}
.kpi .label {{ margin-top:9px; color:var(--text-secondary); font-size:13px; line-height:1.35; }}
.footer {{ position:absolute; left:70px; right:70px; bottom:28px; display:flex; align-items:center; justify-content:space-between; color:rgba(255,255,255,.42); font-size:11px; letter-spacing:.06em; }}
.footer .line {{ position:absolute; left:0; right:0; top:-13px; height:1px; background:linear-gradient(90deg,var(--accent-1),rgba(99,102,241,.35),transparent); }}
.footer .page-no {{ font-family:'Noto Sans Mono','Liberation Mono',monospace; color:var(--accent-1); }}
"""


def shell(title: str, page_number: int, content: str, cover: bool = False) -> str:
    return f"""<!doctype html>
<html lang="de">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width={PAGE_W}, initial-scale=1">
<title>{html.escape(title)}</title>
<style>{common_css()}</style>
</head>
<body>
<div class="grid"></div><div class="scan"></div><div class="corner tl"></div><div class="corner br"></div>
<main class="page">
  <header class="brand">
    <img src="../images/image6.png" alt="FMC3 Robotics">
  </header>
  {content}
  <footer class="footer"><div class="line"></div><span>{FOOTER_TEXT}</span><span class="page-no">0{page_number} / 0{TOTAL_PAGES}</span></footer>
</main>
</body>
</html>
"""


def page_html() -> dict[int, str]:
    feature_items = "".join(
        f'<li><span class="feature-num">{index:02d}</span><span>{html.escape(item)}</span></li>'
        for index, item in enumerate(FEATURES, 1)
    )
    p1 = f"""
<section class="cover-title">
  <div class="eyebrow"><span class="pulse"></span>Einladung · FMC3 Robotics</div>
  <h1>{PAGE_TEXT[1][0]}</h1>
  <p class="subhead">{PAGE_TEXT[1][1]}</p>
</section>
<section class="cover-hero">
  <div class="intro panel">
    <p class="salutation">{PAGE_TEXT[1][2]}</p>
    <p class="body">{PAGE_TEXT[1][3]}</p>
    <p class="body">{PAGE_TEXT[1][4]}</p>
  </div>
  <div class="product image-frame"><img src="../images/image1.png" alt="C5 Reinigungsroboter"></div>
</section>
<section class="feature-section">
  <div class="section-title"><h2>{PAGE_TEXT[1][5]}</h2><div class="rule"></div></div>
  <ul class="features">{feature_items}</ul>
</section>
<section class="spec-row">
  <div class="spec panel"><span class="spec-v">1.980 m²/h</span><span>theoretische Flächenleistung</span></div>
  <div class="spec panel"><span class="spec-v">550 mm</span><span>Schrubbbreite</span></div>
  <div class="spec panel"><span class="spec-v">&lt; 68 dB(A)</span><span>Geräuschpegel</span></div>
  <div class="spec panel"><span class="spec-v">90 L</span><span>großer Wassertank</span></div>
</section>
<style>
.cover-title h1{{max-width:900px;font-size:54px}} .cover-hero{{display:grid;grid-template-columns:1.05fr .95fr;gap:22px;margin-top:28px;align-items:stretch}}
.intro{{padding:24px 26px}} .intro .salutation{{font-size:18px;font-weight:700;margin-bottom:12px}} .intro .body+.body{{margin-top:13px}}
.product{{height:314px;padding:0;background:#11182a}} .product img{{object-fit:cover}}
.feature-section{{margin-top:28px}} .features{{list-style:none;padding:0;margin:0;display:grid;grid-template-columns:1fr 1fr;gap:11px 26px}}
.features li{{min-height:54px;display:grid;grid-template-columns:34px 1fr;gap:8px;align-items:start;color:var(--text-secondary);font-size:14px;line-height:1.43}}
.feature-num{{font-family:'Noto Sans Mono','Liberation Mono',monospace;color:var(--accent-1);font-size:12px;padding-top:2px}}
.spec-row{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-top:24px}} .spec{{padding:16px 12px;text-align:center;min-height:92px}}
.spec-v{{display:block;font-size:24px;font-weight:800;color:var(--text-primary);font-variant-numeric:tabular-nums;margin-bottom:7px}} .spec span:last-child{{font-size:11px;color:var(--text-secondary);line-height:1.3}}
</style>
"""

    p2 = f"""
<div class="eyebrow"><span class="pulse"></span>Business Case · ROI</div>
<div class="section-title"><h1>{PAGE_TEXT[2][0]}</h1><div class="rule"></div></div>
<p class="body lead">{PAGE_TEXT[2][1]}</p>
<section class="kpi-row economics">
  <div class="kpi panel"><div class="value">9.089 €</div><div class="label">Einsparungen pro Jahr</div></div>
  <div class="kpi panel"><div class="value">1,3 Jahre</div><div class="label">Amortisationszeit (ROI)</div></div>
  <div class="kpi panel"><div class="value">1,54 Mio. m²</div><div class="label">gereinigte Fläche pro Jahr</div></div>
</section>
<section class="roi image-frame"><img src="../images/image2.jpeg" alt="Jährlicher Kostenvergleich"></section>
<p class="caption">{PAGE_TEXT[2][8]}</p>
<section class="company">
  <div class="company-copy">
    <div class="section-title"><h2>{PAGE_TEXT[2][9]}</h2><div class="rule"></div></div>
    <p class="body">{PAGE_TEXT[2][10]}</p>
    <p class="body second">{PAGE_TEXT[2][11]}</p>
  </div>
  <aside class="company-facts panel">
    <h3>{PAGE_TEXT[2][12]}</h3>
    <p><strong>Gegründet:</strong> Oktober 2025</p>
    <p><strong>Hauptsitz:</strong> Ingolstadt, Deutschland</p>
    <p><strong>Entwicklung:</strong> Ingolstadt &amp; Shanghai</p>
    <p><strong>Lieferkapazität:</strong> bis zu 30 Roboter/Monat</p>
    <p><strong>Portfolio:</strong> Reinigungs-, Begrüßungs- und Lieferroboter</p>
  </aside>
</section>
<style>
.section-title h1{{font-size:42px;white-space:nowrap}} .lead{{font-size:15px;margin:12px 0 18px}} .economics{{margin-bottom:18px}}
.economics .kpi{{border-top:2px solid var(--accent-1)}} .economics .value{{color:var(--accent-1)}} .roi{{height:535px;padding:10px}} .roi img{{object-fit:contain}}
.company{{display:grid;grid-template-columns:1.55fr .75fr;gap:22px;margin-top:24px}} .company-copy .body{{font-size:14px;line-height:1.52}} .company-copy .second{{margin-top:12px}}
.company-facts{{padding:21px 22px}} .company-facts h3{{color:var(--accent-1);margin-bottom:13px}} .company-facts p{{font-size:13px;color:var(--text-secondary);line-height:1.45;margin-top:8px}} .company-facts strong{{color:var(--text-primary)}}
</style>
"""

    p3 = f"""
<div class="eyebrow"><span class="pulse"></span>People · Product · System</div>
<div class="section-title"><h1>{PAGE_TEXT[3][0]}</h1><div class="rule"></div></div>
<p class="body lead">{PAGE_TEXT[3][1]}</p>
<section class="team image-frame"><img src="../images/image3.jpeg" alt="Das Kernteam von FMC3 Robotics"></section>
<p class="caption">{PAGE_TEXT[3][2]}</p>
<div class="section-title product-title"><h2>{PAGE_TEXT[3][3]}</h2><div class="rule"></div></div>
<p class="body product-lead">{PAGE_TEXT[3][4]}</p>
<section class="parameters image-frame"><img src="../images/image4.png" alt="Technische Daten des C5 und der Arbeitsstation"></section>
<p class="caption">{PAGE_TEXT[3][5]}</p>
<style>
.section-title h1{{font-size:43px}} .lead{{font-size:15px;margin:5px 0 18px}} .team{{height:465px}} .team img{{object-fit:contain}}
.product-title{{margin-top:24px}} .product-lead{{font-size:15px;margin:0 0 16px}} .parameters{{height:530px}} .parameters img{{object-fit:contain}}
</style>
"""

    p4 = f"""
<div class="eyebrow"><span class="pulse"></span>Deployment · Service · Next Step</div>
<div class="section-title"><h1>{PAGE_TEXT[4][0]}</h1><div class="rule"></div></div>
<section class="scenarios image-frame"><img src="../images/image5.png" alt="Anwendungsszenarien, geeignete Bodenbeläge sowie Service- und Supportsystem"></section>
<p class="caption">{PAGE_TEXT[4][1]}</p>
<section class="service-grid">
  <article class="service panel"><div class="service-index">01</div><h3>{PAGE_TEXT[4][2]}</h3><p class="body">{PAGE_TEXT[4][3]}</p></article>
  <article class="service panel"><div class="service-index">02</div><h3>{PAGE_TEXT[4][4]}</h3><p class="body">{PAGE_TEXT[4][5]}</p></article>
</section>
<section class="cta panel">
  <div>
    <div class="eyebrow"><span class="pulse"></span>Live Demonstration</div>
    <h2>{PAGE_TEXT[4][6]}</h2>
  </div>
  <p class="body">{PAGE_TEXT[4][7]}</p>
</section>
<section class="signature">
  <p class="body">{PAGE_TEXT[4][8]}</p>
  <p class="body closing">{PAGE_TEXT[4][9]}</p>
  <div class="name">{PAGE_TEXT[4][10]}</div>
  <div class="role">{PAGE_TEXT[4][11]}</div>
</section>
<style>
.section-title h1{{font-size:39px;white-space:nowrap}} .scenarios{{width:760px;height:690px;margin:18px auto 0}} .scenarios img{{object-fit:contain}}
.service-grid{{display:grid;grid-template-columns:1fr 1fr;gap:18px;margin-top:22px}} .service{{position:relative;padding:22px 24px 20px;min-height:172px}}
.service-index{{position:absolute;right:18px;top:14px;font-family:'Noto Sans Mono','Liberation Mono',monospace;color:rgba(34,211,238,.28);font-size:32px;font-weight:700}}
.service h3{{color:var(--text-primary);margin-bottom:12px;padding-right:45px}} .service .body{{font-size:14px;line-height:1.5}}
.cta{{margin-top:18px;padding:23px 26px;display:grid;grid-template-columns:.72fr 1.28fr;gap:24px;align-items:center;border-color:rgba(34,211,238,.58);box-shadow:0 0 34px rgba(34,211,238,.12)}}
.cta .eyebrow{{font-size:11px;margin-bottom:7px}} .cta h2{{font-size:27px}} .cta .body{{font-size:14px;line-height:1.55}}
.signature{{margin-top:18px;border-left:2px solid var(--accent-1);padding-left:18px}} .signature .body{{font-size:14px}} .signature .closing{{margin-top:8px}}
.name{{margin-top:12px;font-size:18px;font-weight:800}} .role{{margin-top:3px;color:var(--text-secondary);font-size:13px}}
</style>
"""
    return {
        1: shell(PAGE_TEXT[1][0], 1, p1, cover=True),
        2: shell(PAGE_TEXT[2][0], 2, p2),
        3: shell(PAGE_TEXT[3][0], 3, p3),
        4: shell(PAGE_TEXT[4][0], 4, p4),
    }


def build_preview(output_dir: Path) -> None:
    write_text(
        output_dir / "preview.html",
        f"""<!doctype html>
<html lang="de">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>FMC3 C5 Dark Tech Broschüre</title>
<style>
*{{box-sizing:border-box}} body{{margin:0;background:#020611;color:#fff;font-family:'Noto Sans',Arial,sans-serif}}
header{{position:sticky;top:0;z-index:10;height:54px;display:flex;align-items:center;justify-content:space-between;padding:0 24px;background:rgba(5,11,31,.96);border-bottom:1px solid rgba(34,211,238,.25)}}
header strong{{letter-spacing:.04em}} header span{{color:#22D3EE;font-family:monospace}}
main{{padding:28px 16px 60px;display:flex;flex-direction:column;align-items:center;gap:34px}}
.sheet{{width:min(94vw,794px);aspect-ratio:210/297;box-shadow:0 24px 80px rgba(0,0,0,.58);border:1px solid rgba(34,211,238,.2);background:#050b1f;overflow:hidden}}
.page-frame{{width:{PAGE_W}px;height:{PAGE_H}px;border:0;transform-origin:top left}}
</style>
</head>
<body>
<header><strong>FMC3 C5 · Dark Tech Broschüre</strong><span>4 × A4</span></header>
<main>
{''.join(f'<div class="sheet"><iframe class="page-frame" src="slides/slide-{p}.html" title="Seite {p}"></iframe></div>' for p in range(1, TOTAL_PAGES + 1))}
</main>
<script>
function fitPages(){{
  document.querySelectorAll('.sheet').forEach(sheet=>{{
    const frame=sheet.querySelector('.page-frame');
    frame.style.transform=`scale(${{sheet.clientWidth/{PAGE_W}}})`;
  }});
}}
window.addEventListener('resize',fitPages);
fitPages();
</script>
</body>
</html>
""",
    )


def render_pngs(output_dir: Path, scale: int = 2) -> list[Path]:
    if not CHROME:
        raise RuntimeError("google-chrome or chromium is required for HTML screenshots")
    png_dir = output_dir / "png"
    png_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for page in range(1, TOTAL_PAGES + 1):
        html_path = (output_dir / f"slides/slide-{page}.html").resolve()
        png_path = (png_dir / f"slide-{page}.png").resolve()
        command = [
            CHROME,
            "--headless",
            "--no-sandbox",
            "--disable-gpu",
            "--hide-scrollbars",
            "--run-all-compositor-stages-before-draw",
            "--virtual-time-budget=1200",
            f"--force-device-scale-factor={scale}",
            f"--window-size={PAGE_W},{PAGE_H}",
            f"--screenshot={png_path}",
            html_path.as_uri(),
        ]
        completed = subprocess.run(command, capture_output=True, text=True, timeout=60)
        if completed.returncode != 0 or not png_path.exists():
            raise RuntimeError(f"Chrome screenshot failed for page {page}: {completed.stderr}")
        outputs.append(png_path)
    return outputs


def build_pdf(pngs: list[Path], output_path: Path) -> None:
    images = [Image.open(path).convert("RGB") for path in pngs]
    images[0].save(output_path, "PDF", save_all=True, append_images=images[1:], resolution=270)
    for image in images:
        image.close()


def build_pptx(pngs: list[Path], output_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(8.2677165354)
    presentation.slide_height = Inches(11.6929133858)
    blank = presentation.slide_layouts[6]
    for png in pngs:
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(png),
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
    presentation.save(output_path)


def html_text(path: Path) -> str:
    parser = TextExtractor()
    parser.feed(path.read_text(encoding="utf-8"))
    return normalize_text(" ".join(parser.parts))


def content_audit(
    source: Path,
    output_dir: Path,
    source_media: dict[str, dict[str, str | int]],
) -> dict:
    page_texts = {
        page: html_text(output_dir / f"slides/slide-{page}.html")
        for page in range(1, TOTAL_PAGES + 1)
    }
    all_html = " ".join(page_texts.values())
    source_segments = source_text_segments(source)
    missing_source_segments = [
        segment for segment in source_segments
        if normalize_text(segment) not in all_html
    ]
    missing_page_segments = {
        str(page): [
            segment for segment in segments
            if normalize_text(segment) not in page_texts[page]
        ]
        for page, segments in PAGE_TEXT.items()
    }
    missing_page_segments = {
        page: values for page, values in missing_page_segments.items() if values
    }

    media_checks = {}
    for filename, metadata in source_media.items():
        output_file = output_dir / "images" / filename
        output_hash = sha256_bytes(output_file.read_bytes()) if output_file.exists() else None
        media_checks[filename] = {
            "source_sha256": metadata["sha256"],
            "output_sha256": output_hash,
            "byte_identical": output_hash == metadata["sha256"],
        }

    png_checks = {}
    for page in range(1, TOTAL_PAGES + 1):
        png_path = output_dir / f"png/slide-{page}.png"
        with Image.open(png_path) as image:
            png_checks[str(page)] = {
                "dimensions": list(image.size),
                "a4_ratio_ok": abs(image.width / image.height - 210 / 297) < 0.002,
                "bytes": png_path.stat().st_size,
            }

    ok = (
        not missing_source_segments
        and not missing_page_segments
        and all(item["byte_identical"] for item in media_checks.values())
        and all(item["a4_ratio_ok"] and item["bytes"] > 50_000 for item in png_checks.values())
    )
    return {
        "ok": ok,
        "source": str(source),
        "source_text_segment_count": len(source_segments),
        "missing_source_segments": missing_source_segments,
        "missing_page_segments": missing_page_segments,
        "media_checks": media_checks,
        "png_checks": png_checks,
        "note": "Embedded information graphics are preserved byte-identically, retaining all text contained inside those images.",
    }


def write_delivery_manifest(output_dir: Path) -> None:
    manifest = {
        "run_id": output_dir.name,
        "generated_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "summary": {"total_pages": TOTAL_PAGES, "format": "A4 portrait brochure"},
        "artifacts": {
            "preview_html": "preview.html",
            "presentation_png_pptx": "presentation-png.pptx",
            "presentation_svg_pptx": "presentation-svg.pptx",
            "brochure_pdf": "FMC3-C5_Deutsch_Broschure_Business_dark_tech.pdf",
            "content_audit": "content-audit.json",
        },
        "pages": [
            {
                "page": page,
                "planning": f"planning/planning{page}.json",
                "html": f"slides/slide-{page}.html",
                "png": f"png/slide-{page}.png",
            }
            for page in range(1, TOTAL_PAGES + 1)
        ],
        "export_notes": {
            "presentation_png_pptx": "A4 portrait slides with high-resolution page images.",
            "presentation_svg_pptx": "Compatibility fallback: same A4 raster rendering because Node/dom-to-svg is unavailable.",
        },
    }
    write_json(output_dir / "delivery-manifest.json", manifest)


def run_validators(output_dir: Path) -> None:
    commands = [
        [sys.executable, str(ROOT / "scripts/contract_validator.py"), "interview", str(output_dir / "interview-qa.txt")],
        [sys.executable, str(ROOT / "scripts/contract_validator.py"), "requirements-interview", str(output_dir / "requirements-interview.txt")],
        [sys.executable, str(ROOT / "scripts/contract_validator.py"), "source-brief", str(output_dir / "source-brief.txt")],
        [sys.executable, str(ROOT / "scripts/contract_validator.py"), "outline", str(output_dir / "outline.txt")],
        [sys.executable, str(ROOT / "scripts/contract_validator.py"), "style", str(output_dir / "style.json")],
        [sys.executable, str(ROOT / "scripts/planning_validator.py"), str(output_dir / "planning"), "--refs", str(ROOT / "references")],
        [sys.executable, str(ROOT / "scripts/contract_validator.py"), "images", str(output_dir / "planning"), "--require-paths"],
        [sys.executable, str(ROOT / "scripts/contract_validator.py"), "delivery-manifest", str(output_dir / "delivery-manifest.json"), "--base-dir", str(output_dir)],
    ]
    reports = []
    for command in commands:
        completed = subprocess.run(command, cwd=ROOT, capture_output=True, text=True)
        reports.append({
            "command": command,
            "returncode": completed.returncode,
            "stdout": completed.stdout,
            "stderr": completed.stderr,
        })
        if completed.returncode != 0:
            write_json(output_dir / "runtime/validation-report.json", reports)
            raise RuntimeError(f"Validation failed: {' '.join(command)}\n{completed.stdout}\n{completed.stderr}")
    write_json(output_dir / "runtime/validation-report.json", reports)


def build(output_dir: Path, scale: int) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for directory in ("images", "planning", "slides", "png", "runtime"):
        (output_dir / directory).mkdir(parents=True, exist_ok=True)

    media_inventory = extract_media(SOURCE, output_dir / "images")
    workflow_files(output_dir)
    for page, payload in enumerate(planning_payloads(output_dir), 1):
        write_json(output_dir / f"planning/planning{page}.json", payload)
    for page, content in page_html().items():
        write_text(output_dir / f"slides/slide-{page}.html", content)
    build_preview(output_dir)

    pngs = render_pngs(output_dir, scale=scale)
    build_pdf(pngs, output_dir / "FMC3-C5_Deutsch_Broschure_Business_dark_tech.pdf")
    build_pptx(pngs, output_dir / "presentation-png.pptx")
    build_pptx(pngs, output_dir / "presentation-svg.pptx")

    audit = content_audit(SOURCE, output_dir, media_inventory)
    write_json(output_dir / "content-audit.json", audit)
    if not audit["ok"]:
        raise RuntimeError(f"Content audit failed; see {output_dir / 'content-audit.json'}")

    write_delivery_manifest(output_dir)
    run_validators(output_dir)
    print(f"Built dark_tech brochure: {output_dir}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Build the FMC3 C5 dark_tech German brochure")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--scale", type=int, default=2, choices=(1, 2, 3))
    args = parser.parse_args()
    if not SOURCE.exists():
        print(f"Source DOCX not found: {SOURCE}", file=sys.stderr)
        return 1
    build(args.output_dir.resolve(), args.scale)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
